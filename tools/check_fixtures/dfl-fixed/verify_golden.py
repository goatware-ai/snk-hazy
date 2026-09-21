"""verify_golden.py for dfl-freight-audit: re-derive the audit's figures from inputs/ alone."""
import csv, datetime as dt, os, re, sys
from decimal import Decimal as D, ROUND_HALF_UP, ROUND_HALF_EVEN
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

I = HERE / "inputs"
pdate = lambda s: dt.datetime.strptime(s, "%m/%d/%y").date()
BILLS = list(csv.DictReader(open(I / "dfl_freight_bills_q2.csv", newline="")))
BOLS = {r["BOL"]: r for r in csv.DictReader(open(I / "bol_register_q2.csv", newline=""))}
INDEX = {pdate(r["WEEK_OF"]): D(r["MIDWEST_PADD2"]) for r in csv.DictReader(open(I / "doe_diesel_index_2026.csv", newline=""))}
ROWS = [("3.00","18.0"),("3.05","18.0"),("3.10","19.0"),("3.15","19.0"),("3.20","20.0"),("3.25","20.0"),("3.30","20.5"),("3.35","21.5"),("3.40","21.5"),("3.45","22.5"),("3.50","22.5"),("3.55","23.0"),("3.60","24.0"),("3.65","24.0"),("3.70","25.0"),("3.75","25.0"),("3.80","25.5"),("3.85","26.5"),("3.90","26.5"),("3.95","27.5"),("4.00","27.5"),("4.05","28.0"),("4.10","28.5"),("4.15","29.5"),("4.20","30.0"),("4.25","30.0"),("4.30","30.5"),("4.35","31.0"),("4.40","32.0"),("4.45","34.0")]
BANDS = [(D(a), D(b)) for a, b in ROWS]
wb = openpyxl.load_workbook(I / "dfl500_rate_pages.xlsx", data_only=True)
G = ["L5C", "5C", "1M", "2M", "5M"]; GMIN = {"L5C": 0, "5C": 500, "1M": 1000, "2M": 2000, "5M": 5000}
RATES = {(r[0], D(str(r[1]))): {g: D(str(v)) for g, v in zip(G, r[2:7])} for r in wb["DFL 500 2026-A"].iter_rows(min_row=5, values_only=True) if r[0] and r[1] is not None}
ZONES = [(int(a), int(b), r[0]) for r in wb["Zones"].iter_rows(min_row=5, values_only=True) if r[0] for a, b in re.findall(r"(\d{3}) to (\d{3})", r[1])]
INS = {str(r[0]): dict(bw=r[3], cw=r[4], bc=D(str(r[5])), cc=D(str(r[6])), item=r[7]) for r in openpyxl.load_workbook(I / "dfl_weight_inspections_q2.xlsx", data_only=True).active.iter_rows(min_row=5, values_only=True) if r[0]}


def derive(mode="half_up", boundary="begins", unrounded_base=False):
    rnd = ROUND_HALF_EVEN if mode == "half_even" else ROUND_HALF_UP
    r2 = lambda x: D(x).quantize(D("0.01"), rounding=rnd)
    zone = lambda z: next(zz for a, b, zz in ZONES if a <= int(z[:3]) <= b)
    grp = lambda w: "L5C" if w < 500 else "5C" if w < 1000 else "1M" if w < 2000 else "2M" if w < 5000 else "5M"
    def pct(p):
        rows = [v for lo, v in BANDS if (p >= lo if boundary == "begins" else p > lo)]
        return rows[-1] if rows else BANDS[0][1]
    seen, standing, rated_total, claims = set(), {}, D(0), D(0)
    for b in BILLS:
        paid = D(b["TOTAL"])
        if b["PRO"] in seen:
            standing[b["INVOICE"]] = "over"; claims += paid; continue
        seen.add(b["PRO"]); bol = BOLS[b["BOL"]]; w = int(bol["WEIGHT_LB"]); cls = D(bol["NMFC_CLASS"]); wi = False; c = INS.get(b["PRO"])
        if c:
            if abs(c["cw"] - c["bw"]) > D("0.05") * c["bw"]: w = c["cw"]; wi = True
            if c["cc"] != c["bc"] and c["item"]: cls = c["cc"]; wi = True
        rc = D(70) if cls <= 100 else cls; z = zone(bol["DEST_ZIP"]); g = grp(w)
        raw = RATES[(z, rc)][g] * D(w) / 100; base = raw if unrounded_base else r2(raw); ni = G.index(g)
        if ni < 4:
            ng = G[ni + 1]; rawn = RATES[(z, rc)][ng] * D(GMIN[ng]) / 100; nb = rawn if unrounded_base else r2(rawn)
            if nb < base: base = nb
        net = max(r2(base * D("0.32")), D("118.40")); sd = pdate(b["SHIP_DATE"])
        p = pct(INDEX[sd - dt.timedelta(days=sd.weekday())])
        acc = sum([D(85) if bol["LIFTGATE_REQ"] == "Y" else D(0), D(75) if bol["RESIDENTIAL"] == "Y" else D(0), D(22) if bol["NOTIFY_REQ"] == "Y" else D(0), D(28) if wi else D(0)])
        rated = net + r2(net * p / 100) + acc; rated_total += rated; d = paid - rated
        standing[b["INVOICE"]] = "over" if d > 0 else "under" if d < 0 else "correct"
        if d > 0: claims += d
    counts = {k: sum(1 for v in standing.values() if v == k) for k in ("over", "correct", "under")}
    figures = {"rated total": f"{rated_total:,.2f}", "claim total": f"{claims:,.2f}", "overcharged": counts["over"], "billed right": counts["correct"], "under-billed": counts["under"]}
    return figures, standing


EXPECTED = {"rated total": "43,027.40", "claim total": "5,982.91", "overcharged": "71", "billed right": "45", "under-billed": "6"}
VARIANTS = {
    "banker's rounding": ({"mode": "half_even"}, "rounded half up to the cent at each step"),
    "base charge left unrounded": ({"unrounded_base": True}, "rounded half up to the cent at each step"),
    "boundary takes the row that ends there": ({"boundary": "ends"}, "takes the row that begins there"),
}

if __name__ == "__main__":
    rep = Report()
    figures, _ = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures[label], golden)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
