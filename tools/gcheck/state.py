"""Parse-once state for one gate run: the task folder plus every file it has already opened.

Before 2026-09-11 each check opened the workbooks it needed itself, so one gate run on
submissions/22-frankfort-stock-recovery opened workbooks 153 times (the golden 73 times).
`TaskState` is the task folder (it IS a `pathlib.Path`, so every check keeps reading
`folder / "solution"`, `folder.glob(...)` and `p.relative_to(folder)` unchanged) carrying
lazy, memoised parses:

    state.workbook(path, data_only=...)   an openpyxl workbook, one per (path, mode)
    state.document(path)                  a python-docx Document
    state.docx_text(path)                 the document's paragraph text
    state.rows                            the rubric CSV
    state.prompt_text                     prompt.md, "" when absent
    state.metadata                        metadata.json as a dict
    state.input_texts                     [(name, text)] over every readable input
    state.formula_cells                   every formula cell with its cached value
    state.memo(key, fn)                   any other per-run derivation

`core.run_checks` wraps the folder it is given and installs it as `CURRENT` for the
duration of the run, so helpers that receive only a file path (`common.workbook`,
`common._xlsx_prose`, ...) reach the same cache. Outside a run they open the file directly.

Workbooks are always loaded in full (never read_only): a read-only workbook is a stream
that `close()` ends, and every consumer is happy with the full object. A check that
mutates a shared workbook would poison the cache; none does, and none may.
"""
from pathlib import Path

CURRENT = None


def current():
    """The TaskState of the run in progress, or None between runs."""
    return CURRENT


def workbook(path, data_only=False):
    """An openpyxl workbook for `path`, from the running TaskState's cache when there is one."""
    st = CURRENT
    if st is not None:
        return st.workbook(path, data_only=data_only)
    import openpyxl
    return openpyxl.load_workbook(path, data_only=data_only)


def document(path):
    """A python-docx Document for `path`, cached for the run."""
    st = CURRENT
    if st is not None:
        return st.document(path)
    from docx import Document
    return Document(str(path))


class TaskState(Path):
    """The task folder as a Path, plus memoised parses of everything the checks read."""

    def __init__(self, folder, rows=None):
        super().__init__(folder)
        self._cache = {}
        self._rows = rows

    def with_segments(self, *pathsegments):
        # a derived path (folder / "solution") is an ordinary Path; only the task folder
        # itself carries the cache
        return Path(*pathsegments)

    # ---- generic memo -------------------------------------------------------------------

    def memo(self, key, fn):
        """Compute `fn()` once per run under `key` and return it afterwards."""
        if key not in self._cache:
            self._cache[key] = fn()
        return self._cache[key]

    # ---- file parses --------------------------------------------------------------------

    def workbook(self, path, data_only=False):
        key = ("wb", str(path), bool(data_only))
        if key not in self._cache:
            import openpyxl
            self._cache[key] = openpyxl.load_workbook(path, data_only=data_only)
        return self._cache[key]

    def document(self, path):
        key = ("docx", str(path))
        if key not in self._cache:
            from docx import Document
            self._cache[key] = Document(str(path))
        return self._cache[key]

    def docx_text(self, path):
        from .common import _docx_text
        return self.memo(("docx_text", str(path)), lambda: _docx_text(path))

    # ---- the task's parts ---------------------------------------------------------------

    @property
    def rows(self):
        """(NUMBER, CRITERION, WEIGHT) per rubric row, [] when the CSV is absent."""
        if self._rows is None:
            from .common import load_rows, rubric_path
            path = rubric_path(self)
            self._rows = load_rows(path) if path.exists() else []
        return self._rows

    @property
    def prompt_text(self):
        def load():
            p = self / "prompt.md"
            return p.read_text(encoding="utf-8", errors="ignore") if p.is_file() else ""
        return self.memo("prompt_text", load)

    @property
    def metadata(self):
        from .common import load_metadata
        return self.memo("metadata", lambda: load_metadata(self))

    @property
    def input_texts(self):
        """[(file name, readable text)] for every input: csv/txt/md/json bodies, docx text,
        xlsx string cells."""
        from .common import _read_input_texts
        return self.memo("input_texts", lambda: _read_input_texts(self))

    @property
    def formula_cells(self):
        """[(file name, sheet title, coordinate, formula, cached value)] for every formula
        cell of every solution workbook, from one pass over the formula and value modes."""
        from .common import solution_files

        def walk():
            out = []
            for path in solution_files(self, {".xlsx"}):
                try:
                    fwb = self.workbook(path)
                    vwb = self.workbook(path, data_only=True)
                except Exception:
                    continue
                for ws in fwb.worksheets:
                    vs = vwb[ws.title]
                    for row in ws.iter_rows():
                        for c in row:
                            f = c.value
                            if isinstance(f, str) and f.startswith("="):
                                out.append((path.name, ws.title, c.coordinate, f,
                                            vs[c.coordinate].value))
            return out
        return self.memo("formula_cells", walk)
