"""verify_golden.py for evs-staffing-4west: re-derive every figure the staffing plan states from
inputs/ alone under the standards as the golden applies them, and list the alternate readings a
reviewer could take with the phrase the golden uses to settle each one.

    .venv/bin/python drafts/04-evs-staffing-4west/verify_golden.py
"""
import csv
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"


def xr(x, p=2):
    return D(format(D(x), ".15g")).quantize(D(1).scaleb(-p), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(x):,.2f}"


# standards section 1 (minutes), sections 3 to 5
STD = {"occupied": D(22), "terminal": D(43), "isolation": D(58), "nurses": D(15), "med": D(10), "clean": D(12), "soiled": D(12),
       "staff_rr": D(14), "public_rr": D(14), "family": D(20), "conf": D(12), "staff_lounge": D(15), "corridor100": D("1.5"),
       "high_touch": D(30), "burnish": D(90), "scrub": D(360)}
INV = list(csv.DictReader(open(INPUTS / "unit_4west_room_inventory.csv", newline="")))
COUNT = lambda t: sum(1 for r in INV if r["ROOM_TYPE"] == t)
CORR = sum(int(r["SQ_FT"]) for r in INV if r["ROOM_TYPE"] == "Corridor and alcoves")
wb = openpyxl.load_workbook(INPUTS / "census_projection_4west.xlsx", data_only=True)
proj = {r[0]: r for r in wb["Projection"].iter_rows(min_row=5, values_only=True) if r[0]}
AVG = proj["Calendar 2027 average"]; DEC = proj["Dec 2027"]
ws = openpyxl.load_workbook(INPUTS / "evs_roster_wages_oct2026.xlsx", data_only=True)["Wage Scale"]
scale = {r[0]: D(str(r[1])) for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}
RATE = scale["Housekeeper I, base hourly rate"]; DIFF_E = scale["Evening shift differential, per hour worked"]
DIFF_N = scale["Night shift differential, per hour worked"]; DIFF_W = scale["Weekend differential, per hour worked"]
BEN = scale["Employer benefits load on wages and differentials"]; FTE_HOURS = scale["Full-time equivalent, paid hours per year"]
BUDGET = D("205000.00")


def derive(occupied_at_census=False, iso_on_census=False, whole_posts=False, days=365, shifts_per_fte=D("231.4"),
           weekend_all=False, transfers_free=False, basis="average"):
    src = DEC if basis == "peak" else AVG
    adc, dis, iso = D(str(src[1])), D(str(src[2])), D(str(src[3]))
    occ = xr(adc, 2) if occupied_at_census else xr(adc - dis, 2)
    base_iso = adc if iso_on_census else dis
    term_std = xr(dis - base_iso * iso, 3) if iso_on_census else xr(dis * (1 - iso), 3)
    term_iso = xr(base_iso * iso, 3)
    rows = [(occ, STD["occupied"], "occ"), (term_std, STD["terminal"], "term"), (term_iso, STD["isolation"], "term"),
            (D(COUNT("Nurses station")), STD["nurses"], "com"), (D(1), STD["med"], "com"), (D(1), STD["clean"], "com"), (D(2), STD["soiled"], "com"),
            (D(COUNT("Staff restroom")), STD["staff_rr"], "com"), (D(COUNT("Public restroom") * 2), STD["public_rr"], "com"),
            (D(1), STD["family"], "com"), (D(1), STD["conf"], "com"), (D(1), STD["staff_lounge"], "com"), (D(CORR) / 100, STD["corridor100"], "com"),
            (D(2), STD["high_touch"], "com"), (xr(D(1) / 7, 4), STD["burnish"], "floor"), (xr(D(1) / D("30.4"), 4), STD["scrub"], "floor")]
    sh = {"occ": (D(1), D(0), D(0)), "term": (D("0.65"), D("0.30"), D("0.05")), "com": (D("0.5"), D("0.3"), D("0.2")), "floor": (D(0), D(0), D(1))}
    tot = [D(0)] * 4
    for u, m, k in rows:
        d = xr(u * m, 2); tot[0] += d
        for i, s in enumerate(sh[k]):
            tot[i + 1] += xr(d * s, 2)
    posts = [xr(t / 420, 2) for t in tot[1:]]
    if whole_posts:
        posts = [D(int(p) + (1 if p != int(p) else 0)) for p in posts]
    ann = [xr(p * days, 1) for p in posts]
    fte = [xr(a / shifts_per_fte, 2) for a in ann]
    transfers = [D(1), D(1), D(0)]; approved = [D(1), D(0), D(0)]
    hire = [xr(fte[i] - transfers[i] - approved[i], 2) for i in range(3)]

    def cost(f, diff):
        hours = xr(f * FTE_HOURS, 1); base = xr(hours * RATE, 2); sd = xr(hours * diff, 2)
        wk = xr(hours * (D(1) if weekend_all else D(2) / 7) * DIFF_W, 2)
        wd = base + sd + wk; ben = xr(wd * BEN, 2)
        return dict(hours=hours, base=base, sd=sd, wk=wk, wd=wd, ben=ben, total=wd + ben)
    fte_cost = [fte[0] - (transfers[0] if transfers_free else 0), fte[1] - (transfers[1] if transfers_free else 0), fte[2]]
    cs = [cost(fte_cost[0], D(0)), cost(fte_cost[1], DIFF_E), cost(fte_cost[2], DIFF_N)]
    total = sum(c["total"] for c in cs)
    plan3 = cost(D(2), D(0))["total"] + cost(D(1), DIFF_E)["total"]
    figures = {
        "minutes per day": money(tot[0]), "hours per day": f"{xr(tot[0] / 60, 2)}",
        "day minutes": money(tot[1]), "evening minutes": money(tot[2]), "night minutes": money(tot[3]),
        "occupied rooms per day": f"{occ}", "standard terminals per day": f"{term_std}", "isolation terminals per day": f"{term_iso}",
        "day posts": f"{posts[0]}", "evening posts": f"{posts[1]}", "night posts": f"{posts[2]}",
        "day FTE": f"{fte[0]}", "evening FTE": f"{fte[1]}", "night FTE": f"{fte[2]}", "total FTE": f"{sum(fte)}",
        "still to hire": f"{sum(hire)}", "evening surplus": f"{-hire[1]}",
        "day cost": money(cs[0]["total"]), "evening cost": money(cs[1]["total"]), "night cost": money(cs[2]["total"]),
        "total cost": money(total), "under budget": money(BUDGET - total), "working plan cost": money(plan3),
        "day weekend differential": money(cs[0]["wk"]), "day benefits load": money(cs[0]["ben"]),
    }
    standings = {"day FTE": figures["day FTE"], "evening FTE": figures["evening FTE"], "night FTE": figures["night FTE"],
                 "total FTE": figures["total FTE"], "total cost": figures["total cost"],
                 "budget": "under" if total <= BUDGET else "over", "plan holds": "no" if sum(fte) > D(3) else "yes"}
    return figures, standings


VARIANTS = {
    "occupied cleans counted at the full census": ({"occupied_at_census": True}, "the census less discharges"),
    "isolation share applied to the census": ({"iso_on_census": True}, "split by the isolation share of discharges"),
    "posts rounded up to whole people": ({"whole_posts": True}, "carried to two decimals"),
    "five-day week": ({"days": 260}, "365 shifts a year"),
    "no non-productive allowance": ({"shifts_per_fte": D(260)}, "231.4 worked shifts"),
    "weekend differential on every paid hour": ({"weekend_all": True}, "two-sevenths of paid hours"),
    "transferred staff costed at zero": ({"transfers_free": True}, "transfers included"),
    "December peak as the plan basis": ({"basis": "peak"}, "built on the calendar 2027 average"),
}

EXPECTED = {
    "minutes per day": "1,019.84", "hours per day": "17.00", "day minutes": "698.98", "evening minutes": "211.42", "night minutes": "109.44",
    "occupied rooms per day": "13.20", "standard terminals per day": "6.724", "isolation terminals per day": "1.476",
    "day posts": "1.66", "evening posts": "0.50", "night posts": "0.26",
    "day FTE": "2.62", "evening FTE": "0.79", "night FTE": "0.41", "total FTE": "3.82", "still to hire": "0.82", "evening surplus": "0.21",
    "day cost": "126,011.30", "evening cost": "40,300.37", "night cost": "21,622.14", "total cost": "187,933.81",
    "under budget": "17,066.19", "working plan cost": "147,204.95", "day weekend differential": "1,557.03", "day benefits load": "27,178.91",
}

if __name__ == "__main__":
    rep = Report()
    figures, standings = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures.get(label), golden)
    peak, _ = derive(basis="peak")
    rep.expect("December peak total FTE", peak["total FTE"], "4.04")
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
