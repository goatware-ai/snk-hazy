"""verify_golden.py for tract7-boundary-retracement: re-derive every figure the analysis states from
inputs/ alone under the firm's standards as the golden applies them, and list the alternate readings
a reviewer could take with the phrase the golden uses to settle each one.

    .venv/bin/python drafts/06-tract7-boundary-retracement/verify_golden.py
"""
import csv
import math
import os
import re
import sys
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
GOLDEN = HERE / "solution" / "tract7_boundary_analysis.xlsx"
POB = (5000.0, 5000.0); RURAL = 10000; URBAN = 15000; SQFT_ACRE = 43560; DEED_THR = 1.0; CALL_TOL = 1.0


def azimuth(ns, ang, ew):
    k = (ns, ew)
    return ang if k == ("N", "E") else 180 - ang if k == ("S", "E") else 180 + ang if k == ("S", "W") else 360 - ang


def quad(az):
    if az <= 90: return "N", az, "E"
    if az <= 180: return "S", 180 - az, "E"
    if az <= 270: return "S", az - 180, "W"
    return "N", 360 - az, "W"


def dms(ang):
    d = int(ang); m = int((ang - d) * 60); s = round(((ang - d) * 60 - m) * 60)
    if s == 60: s = 0; m += 1
    if m == 60: m = 0; d += 1
    return d, m, s


def bearing_text(az):
    ns, ang, ew = quad(az); d, m, s = dms(ang)
    return f"{ns} {d}° {m:02d}' {s:02d}\" {ew}"


# ---------------------------------------------------------------- inputs
shots = list(csv.DictReader(open(INPUTS / "field_notes_traverse_tract7.csv", newline="")))
courses = {}
for s in shots:
    az = azimuth(s["NS"], int(s["DEG"]) + int(s["MIN"]) / 60 + int(s["SEC"]) / 3600, s["EW"])
    courses.setdefault(s["COURSE"], {})[s["DIRECTION"]] = (az, float(s["HORIZ_DIST_FT"]))
ORDER = list(courses)

deed_doc = Document(INPUTS / "deed_book_412_page_233.docx")
deed_text = "\n".join(p.text for p in deed_doc.paragraphs)
CALL = re.compile(r"([NS]) (\d+)° (\d+)' ([EW]) (\d+\.\d) feet")
deed_calls = [(ns, azimuth(ns, int(d) + int(m) / 60, ew), float(dist), int(d), int(m), ew) for ns, d, m, ew, dist in CALL.findall(deed_text)]
assert len(deed_calls) == 5, deed_calls
deed_acres = float(re.search(r"containing (\d+\.\d+) acres", deed_text).group(1))
std_text = "\n".join(p.text for p in Document(INPUTS / "firm_survey_standards.docx").paragraphs)
assert "one part in 10,000" in std_text and "43,560" in std_text


def derive(rotate=True, adjust=False, standard=RURAL, lat_places=3):
    fig, standing = {}, {}
    # traverse reduction: mean of forward and reverse (reverse turned 180)
    mean = []
    for c in ORDER:
        fa, fd = courses[c]["Forward"]; ra, rd = courses[c]["Reverse"]
        az = (fa + (ra + 180) % 360) / 2; dist = (fd + rd) / 2
        lat = round(dist * math.cos(math.radians(az)), lat_places); dep = round(dist * math.sin(math.radians(az)), lat_places)
        mean.append((c, az, dist, lat, dep))
    per = sum(m[2] for m in mean); sl = round(sum(m[3] for m in mean), 3); sd = round(sum(m[4] for m in mean), 3)
    if adjust:   # compass rule closes the figure exactly, so the closure reported would be zero
        mean = [(c, az, d, lat - sl * d / per, dep - sd * d / per) for c, az, d, lat, dep in mean]
        sl = round(sum(m[3] for m in mean), 3); sd = round(sum(m[4] for m in mean), 3)
    err = round(math.sqrt(sl ** 2 + sd ** 2), 3)
    prec = round(per / err) if err else None
    fig["perimeter"] = f"{per:.3f}"; fig["lat misclosure"] = f"{sl:.3f}"; fig["dep misclosure"] = f"{sd:.3f}"
    fig["linear error"] = f"{err:.3f}"; fig["precision"] = prec
    standing["closure test"] = "Meets the standard" if prec is not None and prec >= standard else "Does not meet the standard"
    fig["closure test"] = standing["closure test"]
    fig["mean bearing 1-2"] = bearing_text(mean[0][1])
    fig["mean distance 3-4"] = f"{mean[2][2]:.3f}"
    # coordinates and area
    N, E = [POB[0]], [POB[1]]
    for c, az, d, lat, dep in mean[:-1]:
        N.append(N[-1] + lat); E.append(E[-1] + dep)
    area = round(0.5 * abs(sum(N[i] * E[(i + 1) % 5] - N[(i + 1) % 5] * E[i] for i in range(5))))
    acres = round(area / SQFT_ACRE, 3)
    fig["area sqft"] = area; fig["acres"] = f"{acres:.3f}"; fig["deed acres"] = f"{deed_acres:.2f}"; fig["acre difference"] = f"{round(acres - deed_acres, 3):.3f}"
    for i in range(5):
        fig[f"corner {i + 1} N"] = f"{N[i]:.3f}"; fig[f"corner {i + 1} E"] = f"{E[i]:.3f}"
    # deed on its own calls
    dl = round(sum(round(d * math.cos(math.radians(az)), 3) for _, az, d, *_ in deed_calls), 3)
    dd = round(sum(round(d * math.sin(math.radians(az)), 3) for _, az, d, *_ in deed_calls), 3)
    dper = round(sum(d for _, _, d, *_ in deed_calls), 1); derr = round(math.sqrt(dl ** 2 + dd ** 2), 3)
    per1000 = round(derr / dper * 1000, 3)
    fig["deed perimeter"] = f"{dper:.1f}"; fig["deed misclosure"] = f"{derr:.3f}"; fig["deed per 1000"] = f"{per1000:.3f}"; fig["deed precision"] = round(dper / derr)
    standing["deed test"] = "Defect in the description" if per1000 > DEED_THR else "Within the tolerance of the original survey"
    fig["deed test"] = standing["deed test"]
    rot = mean[0][1] - deed_calls[0][1] if rotate else 0.0
    rd_, rm_ = int(rot), round((rot - int(rot)) * 60)
    fig["rotation text"] = f"{rd_}° {rm_:02d}' east of the deed bearings"
    flags = []
    for i in range(5):
        raz = (deed_calls[i][1] + rot) % 360
        bdiff = round((mean[i][1] - raz) * 60, 2); ddiff = round(mean[i][2] - deed_calls[i][2], 2)
        fig[f"bearing diff {ORDER[i]}"] = f"{bdiff:.2f}"; fig[f"distance diff {ORDER[i]}"] = f"{ddiff:.2f}"
        flags.append(abs(ddiff) > CALL_TOL)
    standing["calls over tolerance"] = sum(flags)
    standing["call in error"] = ORDER[flags.index(True)] if any(flags) else "none"
    standing["bearings within 2 minutes after rotation"] = sum(1 for i in range(5) if abs(float(fig[f"bearing diff {ORDER[i]}"])) <= 2)
    fig["calls over tolerance"] = standing["calls over tolerance"]; fig["call in error"] = standing["call in error"]
    fig["deed distance of call in error"] = f"{deed_calls[flags.index(True)][2]:.1f}" if any(flags) else "none"
    return fig, standing


VARIANTS = {
    "compass rule applied before the closure is reported": ({"adjust": True}, "unadjusted"),
    "deed bearings compared without rotation": ({"rotate": False}, "rotated by that amount before comparison"),
    "urban closure standard applied": ({"standard": URBAN}, "rural"),
    "latitudes carried to two decimals": ({"lat_places": 2}, "three decimal"),
}

fig, standing = derive()
wb = openpyxl.load_workbook(GOLDEN, data_only=True)
g = wb["Findings"]; tr = wb["Traverse"]; de = wb["Deed"]; ar = wb["Area"]; mo = wb["Monuments"]
rep = Report()
rep.expect("perimeter", fig["perimeter"], f"{g['B7'].value:.3f}")
rep.expect("lat misclosure", fig["lat misclosure"], f"{g['B8'].value:.3f}")
rep.expect("dep misclosure", fig["dep misclosure"], f"{g['B9'].value:.3f}")
rep.expect("linear error", fig["linear error"], f"{g['B10'].value:.3f}")
rep.expect("precision", fig["precision"], g["B11"].value)
rep.expect("closure test", fig["closure test"], g["B13"].value)
rep.expect("mean bearing 1-2", fig["mean bearing 1-2"], tr["G4"].value)
rep.expect("mean distance 3-4", fig["mean distance 3-4"], f"{tr['J6'].value:.3f}")
rep.expect("deed misclosure", fig["deed misclosure"], f"{g['B15'].value:.3f}")
rep.expect("deed per 1000", fig["deed per 1000"], f"{g['B16'].value:.3f}")
rep.expect("deed perimeter", fig["deed perimeter"], f"{de['B12'].value:.1f}")
rep.expect("deed precision", fig["deed precision"], de["B17"].value)
rep.expect("deed test", fig["deed test"], g["B17"].value)
rep.expect("rotation text", fig["rotation text"], g["B18"].value)
rep.expect("calls over tolerance", fig["calls over tolerance"], g["B19"].value)
rep.expect("call in error", fig["call in error"], g["B20"].value)
rep.expect("deed distance of call in error", fig["deed distance of call in error"], f"{g['B21'].value:.1f}")
rep.expect("measured distance of call in error", fig["mean distance 3-4"], f"{g['B22'].value:.3f}")
for i in range(5):
    r = 4 + i
    rep.expect(f"bearing diff {ORDER[i]}", fig[f"bearing diff {ORDER[i]}"], f"{de[f'N{r}'].value:.2f}")
    rep.expect(f"distance diff {ORDER[i]}", fig[f"distance diff {ORDER[i]}"], f"{de[f'O{r}'].value:.2f}")
    rep.expect(f"corner {i + 1} N", fig[f"corner {i + 1} N"], f"{ar[f'B{r}'].value:.3f}")
    rep.expect(f"corner {i + 1} E", fig[f"corner {i + 1} E"], f"{ar[f'C{r}'].value:.3f}")
rep.expect("area sqft", fig["area sqft"], g["B25"].value)
rep.expect("acres", fig["acres"], f"{g['B26'].value:.3f}")
rep.expect("deed acres", fig["deed acres"], f"{g['B27'].value:.2f}")
rep.expect("acre difference", fig["acre difference"], f"{g['B28'].value:.3f}")
rep.expect("monuments found", 4, g["B31"].value)
rep.expect("monuments set", 1, g["B32"].value)
rep.expect("corner set", 3, g["B33"].value)
# the note's figures
note = "\n".join(str(c.value) for row in wb["Note to Ellen"].iter_rows() for c in row if c.value)
for phrase in (f"one part in {fig['precision']:,}", f"{fig['linear error']} feet", f"{fig['deed misclosure']} feet", f"{fig['area sqft']:,} square feet",
               f"{fig['acres']} acres", "2 degrees 16 minutes", f"{fig['deed distance of call in error']} feet", f"{float(fig['mean distance 3-4']):.2f} feet"):
    rep.expect(f"note states {phrase}", phrase in note, True)
rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
rep.finish()
