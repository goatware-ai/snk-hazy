"""verify_golden.py for covenant-certificate-aug-2026: re-derive every figure the certificate
states from inputs/ alone, under the conventions the golden's own text states, and list every
alternate reading a reviewer could take with the phrase the golden uses to settle it.

    .venv/bin/python drafts/01-covenant-certificate-aug-2026/verify_golden.py
"""
import csv
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"


def xround(x, places=2):
    """Excel ROUND: half up off the 15 significant digit value."""
    return D(format(D(x), ".15g")).quantize(D(1).scaleb(-places), rounding=ROUND_HALF_UP)


def money(s):
    return f"{D(s):,.2f}"


# ------------------------------------------------------------------ ledger export
GL = {}
with open(INPUTS / "gl_activity_sep25_aug26.csv", newline="") as fh:
    for row in csv.DictReader(fh):
        months = [k for k in row if k not in ("ACCOUNT", "DESCRIPTION", "STATEMENT")]
        GL[int(row["ACCOUNT"])] = sum((D(row[m]) for m in months), D(0))

PL = [a for a in GL if a != 3900]
CASH_INT = (7000, 7005, 7010, 7015)
NONCASH_INT = (7018,)
INT_INCOME = (7050,)
TAX = (8000, 8010)
DA = (5050, 6090, 6095)
WRITEOFF = (7030,)
GAIN = (7020,)
SEVERANCE = 6100
SETTLEMENT = 6110
DIVIDEND = 3900

# ------------------------------------------------------------------ treasury schedules
wb = openpyxl.load_workbook(INPUTS / "treasury_schedules_083126.xlsx", data_only=True)


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


tl = rows(wb["Term Loan"])
TL_BAL = next(D(str(r[3])) for r in tl if r[0] == "08/31/2026" and r[1] == "Scheduled principal payment")
TL_SCHED = -sum((D(str(r[2])) for r in tl if r[1] == "Scheduled principal payment"
                and r[0] in ("11/30/2025", "02/28/2026", "05/31/2026", "08/31/2026")), D(0))
TL_PREPAY = -sum((D(str(r[2])) for r in tl if r[1] == "Voluntary prepayment"), D(0))
rv = rows(wb["Revolver"])
REV_BAL = next(D(str(r[1])) for r in rv if r[0] == "08/31/2026")
sn = rows(wb["Seller Note"])
SN_BAL = next(D(str(r[4])) for r in sn if r[0] == "06/01/2026")
SN_SCHED = -next(D(str(r[2])) for r in sn if r[0] == "06/01/2026")
ls = rows(wb["Leases"])
LEASES = [dict(id=r[0], cls=r[3], bal=D(str(r[7])), prin=D(str(r[8]))) for r in ls
          if isinstance(r[0], str) and r[0][:3] in ("FL-", "OL-") and r[3] in ("Finance", "Operating")]
lc = rows(wb["LCs"])
LC_FACE = sum((D(str(r[5])) for r in lc if isinstance(r[0], str) and r[0].startswith("SBLC-")), D(0))
cx = rows(wb["Capex"])
CAPEX = [(r[0], D(str(r[3])), r[4]) for r in cx if isinstance(r[0], str) and len(r[0]) == 6 and r[0][3] == "-" and r[3] is not None]
tx = rows(wb["Income Taxes"])
TAXPAY = [(r[0], D(str(r[2])), D(str(r[3]))) for r in tx
          if isinstance(r[0], str) and r[0].count("/") == 2 and r[2] is not None]

# ------------------------------------------------------------------ Ostrander statements
ob = openpyxl.load_workbook(INPUTS / "ostrander_pl_sep25_feb26.xlsx", data_only=True)
OSH = {}
for r in rows(ob["Income Statement"]):
    if isinstance(r[0], str) and r[0].startswith("   "):
        OSH[r[0].strip()] = sum((D(str(v)) for v in r[1:7] if v is not None), D(0))
O_INCOME = ("Sales", "Sales returns")


def derive(settlement=False, sev_uncapped=False, writeoff_under_e=False, cap_excl_acq=False,
           lcs_out=False, oper_in=False, prepay_sched=False, dfc_in_cash=False, no_net_income=False,
           auto_added=False, owner_not_added=False, max_300=False, fin_capex_in=False,
           refund_not_netted=False, no_dividend=False):
    net_income = -sum((GL[a] for a in PL), D(0))
    int_cash = sum((GL[a] for a in CASH_INT), D(0))
    int_noncash = sum((GL[a] for a in NONCASH_INT), D(0))
    int_income = sum((GL[a] for a in INT_INCOME), D(0))
    int_net = int_cash + int_noncash + (D(0) if no_net_income else int_income)
    taxes = sum((GL[a] for a in TAX), D(0))
    da = sum((GL[a] for a in DA), D(0))
    writeoff = sum((GL[a] for a in WRITEOFF), D(0))
    gain = sum((GL[a] for a in GAIN), D(0))
    o_ni = sum((OSH[k] for k in OSH if k in O_INCOME), D(0)) - sum((OSH[k] for k in OSH if k not in O_INCOME), D(0))
    acq = o_ni + OSH["Interest expense - line of credit"] + OSH["Depreciation"]
    if not owner_not_added:
        acq += OSH["Officer salary - D. Ostrander"]
    if auto_added:
        acq += OSH["Auto and travel - D. Ostrander"]
    before_e = net_income + int_net + taxes + da + gain + acq + (D(0) if writeoff_under_e else writeoff)
    cap_base = before_e - acq if cap_excl_acq else before_e
    cap = xround(cap_base * D("0.10"))
    requested = GL[SEVERANCE] if sev_uncapped else min(GL[SEVERANCE], D("412000.00"))
    if settlement:
        requested += GL[SETTLEMENT]
    if writeoff_under_e:
        requested += writeoff
    permitted = min(requested, cap)
    ebitda = before_e + permitted
    fl_bal = sum((L["bal"] for L in LEASES if L["cls"] == "Finance"), D(0))
    ol_bal = sum((L["bal"] for L in LEASES if L["cls"] == "Operating"), D(0))
    funded = TL_BAL + REV_BAL + SN_BAL + fl_bal + (D(0) if lcs_out else LC_FACE) + (ol_bal if oper_in else D(0))
    leverage = xround(funded / ebitda)
    lev_max = D("3.00") if max_300 else D("2.75")
    cash_int = int_cash + (int_noncash if dfc_in_cash else D(0)) + (D(0) if no_net_income else int_income)
    fl_prin = sum((L["prin"] for L in LEASES if L["cls"] == "Finance"), D(0))
    sched = TL_SCHED + SN_SCHED + fl_prin + (TL_PREPAY if prepay_sched else D(0))
    fixed = cash_int + sched
    capex_total = sum((a for _, a, _ in CAPEX), D(0))
    capex_fin = sum((a for _, a, f in CAPEX if f == "Finance lease"), D(0))
    unfin = capex_total - (D(0) if fin_capex_in else capex_fin)
    cash_taxes = sum((f + s for _, f, s in TAXPAY if not (refund_not_netted and f < 0)), D(0))
    dividend = D(0) if no_dividend else GL[DIVIDEND]
    numerator = ebitda - unfin - cash_taxes - dividend
    fccr = xround(numerator / fixed)
    required = xround(fixed * D("1.20"))
    cure = required - numerator
    room = cap - permitted
    sens_num = numerator + room
    sens_fccr = xround(sens_num / fixed)
    capex_fy26 = sum((a for m, a, _ in CAPEX if m not in ("Sep-25", "Oct-25", "Nov-25")), D(0))
    figures = {
        "Consolidated Net Income": money(net_income),
        "Consolidated Interest Expense, net": money(int_net),
        "Provision for income taxes": money(taxes),
        "Depreciation and amortization": money(da),
        "Acquired EBITDA": money(acq),
        "Consolidated EBITDA before clause (e)": money(before_e),
        "Clause (e) cap": money(cap),
        "Clause (e) permitted": money(permitted),
        "Consolidated EBITDA": money(ebitda),
        "Severance above the approved amount": money(GL[SEVERANCE] - min(GL[SEVERANCE], D("412000.00"))),
        "Consolidated Funded Debt": money(funded),
        "Finance lease obligations": money(fl_bal),
        "Total Leverage Ratio": f"{leverage}",
        "Consolidated Cash Interest Expense": money(cash_int),
        "Scheduled Principal Payments": money(sched),
        "Fixed Charges": money(fixed),
        "Capital Expenditures, Test Period": money(capex_total),
        "Unfinanced Capital Expenditures": money(unfin),
        "Capital Expenditures, FY2026 to date": money(capex_fy26),
        "Cash Taxes": money(cash_taxes),
        "Restricted Payments": money(dividend),
        "Numerator": money(numerator),
        "Fixed Charge Coverage Ratio": f"{fccr}",
        "Numerator required": money(required),
        "Cure Amount": money(cure),
        "Room under the clause (e) cap": money(room),
        "FCCR with the settlement approved": f"{sens_fccr}",
        "Ostrander officer salary": money(OSH["Officer salary - D. Ostrander"]),
        "Ostrander auto and travel, not added": money(OSH["Auto and travel - D. Ostrander"]),
    }
    standings = {
        "Section 6.11": "in compliance" if leverage <= lev_max else "not in compliance",
        "Section 6.12": "in compliance" if fccr >= D("1.20") else "not in compliance",
        "Total Leverage Ratio": f"{leverage}",
        "Fixed Charge Coverage Ratio": f"{fccr}",
        "Consolidated EBITDA": money(ebitda),
        "Consolidated Funded Debt": money(funded),
        "Fixed Charges": money(fixed),
        "Cure Amount": money(cure),
        "Numerator": money(numerator),
        "FCCR with the settlement approved": f"{sens_fccr}",
    }
    return figures, standings


# Each variant: name -> (keyword overrides for derive, the phrase the golden states to settle it)
VARIANTS = {
    "Kowalczyk settlement added back under clause (e)": ({"settlement": True}, "is not in the add-backs"),
    "severance added back at the full expense": ({"sev_uncapped": True}, "added back at $412,000.00"),
    "leasehold write-off counted under clause (e)": ({"writeoff_under_e": True}, "does not count against the clause (e) cap"),
    "clause (e) cap base without Acquired EBITDA": ({"cap_excl_acq": True}, "of Consolidated EBITDA before clause (e)"),
    "letters of credit left out of Funded Debt": ({"lcs_out": True}, "count whether or not drawn"),
    "operating lease liabilities in Funded Debt": ({"oper_in": True}, "outside the definition"),
    "April 15 prepayment counted as scheduled principal": ({"prepay_sched": True}, "was voluntary and is excluded"),
    "deferred financing amortization in cash interest": ({"dfc_in_cash": True}, "without the amortization of deferred financing costs"),
    "interest income not netted": ({"no_net_income": True}, "net of interest income"),
    "Ostrander auto and travel added back": ({"auto_added": True}, "are not added"),
    "Ostrander officer salary not added back": ({"owner_not_added": True}, "the one adjustment approved"),
    "leverage tested at 3.00 to 1.00": ({"max_300": True}, "steps the maximum down to 2.75 to 1.00"),
    "Brandt-financed center counted as unfinanced": ({"fin_capex_in": True}, "financed on the Brandt lease"),
    "January refund not netted from Cash Taxes": ({"refund_not_netted": True}, "net of the $42,610.00 federal refund"),
    "dividend left out of Restricted Payments": ({"no_dividend": True}, "as a Restricted Payment"),
}

# Each figure the golden states: label -> the text the golden prints
EXPECTED = {
    "Consolidated Net Income": "1,175,708.86",
    "Consolidated Interest Expense, net": "1,073,912.26",
    "Provision for income taxes": "535,462.28",
    "Depreciation and amortization": "1,808,911.12",
    "Acquired EBITDA": "434,009.02",
    "Consolidated EBITDA before clause (e)": "5,002,830.01",
    "Clause (e) cap": "500,283.00",
    "Clause (e) permitted": "412,000.00",
    "Consolidated EBITDA": "5,414,830.01",
    "Severance above the approved amount": "26,912.55",
    "Consolidated Funded Debt": "14,167,103.23",
    "Finance lease obligations": "1,029,603.23",
    "Total Leverage Ratio": "2.62",
    "Consolidated Cash Interest Expense": "1,015,912.30",
    "Scheduled Principal Payments": "1,752,783.47",
    "Fixed Charges": "2,768,695.77",
    "Capital Expenditures, Test Period": "1,801,015.83",
    "Unfinanced Capital Expenditures": "1,188,565.83",
    "Capital Expenditures, FY2026 to date": "1,420,452.78",
    "Cash Taxes": "583,845.00",
    "Restricted Payments": "437,500.00",
    "Numerator": "3,204,919.18",
    "Fixed Charge Coverage Ratio": "1.16",
    "Numerator required": "3,322,434.92",
    "Cure Amount": "117,515.74",
    "Room under the clause (e) cap": "88,283.00",
    "FCCR with the settlement approved": "1.19",
    "Ostrander officer salary": "111,000.00",
    "Ostrander auto and travel, not added": "9,225.49",
}

if __name__ == "__main__":
    rep = Report()
    figures, _ = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures.get(label), golden)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
