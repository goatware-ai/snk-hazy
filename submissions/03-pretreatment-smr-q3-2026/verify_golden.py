"""verify_golden.py for pretreatment-smr-q3-2026: re-derive every figure and standing the report states
from inputs/ alone under the permit's computation rules as the golden states them, and list the alternate
readings a reviewer could take with the phrase the golden uses to settle each one.

    .venv/bin/python drafts/03-pretreatment-smr-q3-2026/verify_golden.py
"""
import datetime as dt
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
PARAMS = ["Cd", "Cr", "Cu", "Ni", "Zn", "CN"]
DAILY = {"Cd": D("0.11"), "Cr": D("2.77"), "Cu": D("3.38"), "Ni": D("3.98"), "Zn": D("2.61"), "CN": D("1.20")}
MONTHLY = {"Cd": D("0.07"), "Cr": D("1.71"), "Cu": D("2.07"), "Ni": D("2.38"), "Zn": D("1.48"), "CN": D("0.65")}
FLOW_LIMIT = 45000


def xround(x, places):
    return D(format(D(x), ".15g")).quantize(D(1).scaleb(-places), rounding=ROUND_HALF_UP)


wb = openpyxl.load_workbook(INPUTS / "lab_results_apr_sep2026.xlsx", data_only=True)
ws = wb["Results"]
rows = list(ws.iter_rows(values_only=True))
hdr = [str(h) for h in rows[3]]
EVENTS = []
for r in rows[4:]:
    if not (isinstance(r[0], str) and r[0].startswith("VP-")):
        continue
    rec = dict(id=r[0], date=dt.datetime.strptime(r[1], "%m/%d/%Y").date(), kind=r[3], report=dt.datetime.strptime(r[hdr.index("REPORT DATE")], "%m/%d/%Y").date())
    for p in PARAMS:
        rec[p] = (r[hdr.index(f"{p} MG/L")], (r[hdr.index(f"{p} QUAL")] or "").strip())
    EVENTS.append(rec)
mdl = {}
for r in wb["Methods"].iter_rows(min_row=4, values_only=True):
    if isinstance(r[0], str) and r[0] != "pH":
        key = {"Cadmium, total": "Cd", "Chromium, total": "Cr", "Copper, total": "Cu", "Nickel, total": "Ni", "Zinc, total": "Zn", "Cyanide, total": "CN"}[r[0]]
        mdl[key] = D(str(r[2]))

log = Document(INPUTS / "operator_log_sep2026.docx")
FLOWS = {}
for t in log.tables:
    for row in t.rows[1:]:
        cells = [c.text for c in row.cells]
        if cells[0][:2] == "09":
            FLOWS[int(cells[0][3:5])] = int(cells[1].replace(",", ""))
EVENT_TIME = dt.datetime(2026, 9, 9, 7, 50)
NOTICE_TIME = dt.datetime(2026, 9, 10, 16, 30)


def derive(nondetect="zero", include_invalid=False, resample_in_month=True, resample_in_snc=True, trc="1.2",
           notice_from_manager=False):
    def treated(p, rec):
        val, q = rec[p]
        if q == "NS":
            return None
        if q == "H" and not include_invalid:
            return None
        if q == "U":
            return {"zero": D(0), "mdl": mdl[p], "half": mdl[p] / 2}[nondetect]
        return D(str(val))
    figures, standings = {}, {}
    for m, label in ((7, "Jul"), (8, "Aug"), (9, "Sep")):
        for p in PARAMS:
            vals = [treated(p, e) for e in EVENTS if e["date"].month == m and (resample_in_month or e["kind"] != "Resample")]
            vals = [v for v in vals if v is not None]
            n = len(vals)
            avg = xround(sum(vals) / n, 3) if n else None
            mx = max(vals) if n else None
            figures[f"{label} {p} valid samples"] = str(n)
            figures[f"{label} {p} monthly average"] = f"{avg:.3f}" if avg is not None else "none"
            figures[f"{label} {p} daily maximum"] = f"{mx:.3f}" if mx is not None else "none"
            standings[f"{label} {p} monthly average status"] = "Met" if avg is not None and avg <= MONTHLY[p] else "Exceeded"
            standings[f"{label} {p} daily maximum status"] = "Met" if mx is not None and mx <= DAILY[p] else "Exceeded"
            standings[f"{label} {p} frequency status"] = "Met" if n >= 2 else "Below frequency"
    for p in PARAMS:
        vals = [treated(p, e) for e in EVENTS if resample_in_snc or e["kind"] != "Resample"]
        vals = [v for v in vals if v is not None]
        n = len(vals)
        thr = xround(DAILY[p] * D(trc), 3)
        above = sum(1 for v in vals if v > DAILY[p]); above_trc = sum(1 for v in vals if v > thr)
        share = xround(D(above) / n * 100, 1); share_trc = xround(D(above_trc) / n * 100, 1)
        figures[f"SNC {p} measurements"] = str(n); figures[f"SNC {p} above limit"] = str(above)
        figures[f"SNC {p} above TRC"] = str(above_trc); figures[f"SNC {p} share above TRC pct"] = f"{share_trc}"
        figures[f"SNC {p} TRC threshold"] = f"{thr:.3f}"
        standings[f"SNC {p}"] = "Significant noncompliance" if share >= 66 or share_trc >= 33 else "Not in significant noncompliance"
    figures["September highest daily flow"] = str(max(FLOWS.values()))
    figures["September days above the flow limit"] = str(sum(1 for v in FLOWS.values() if v > FLOW_LIMIT))
    figures["September flow total"] = str(sum(FLOWS.values()))
    start = dt.datetime(2026, 9, 9, 8, 30) if notice_from_manager else EVENT_TIME
    late_by = NOTICE_TIME - (start + dt.timedelta(hours=24))
    standings["overflow telephone notice"] = "Late" if late_by.total_seconds() > 0 else "Met"
    figures["hours from the overflow to the telephone notice"] = f"{(NOTICE_TIME - EVENT_TIME).total_seconds() / 3600:.2f}"
    # resample deadlines: 30 days after the lab report of each daily maximum exceedance
    for e in EVENTS:
        for p in PARAMS:
            v = treated(p, e)
            if v is not None and v > DAILY[p]:
                due = e["report"] + dt.timedelta(days=30)
                later = [x for x in EVENTS if x["date"] > e["date"] and x["date"] <= due and treated(p, x) is not None]
                standings[f"resample after {e['id']} {p}"] = "Met" if later else "Open"
                figures[f"resample due after {e['id']} {p}"] = f"{due:%m/%d/%Y}"
    figures["September Cu average without the resample"] = f"{xround(sum(D(str(e['Cu'][0])) for e in EVENTS if e['date'].month == 9 and e['kind'] == 'Routine') / 2, 3):.3f}"
    return figures, standings


VARIANTS = {
    "non-detects carried at the detection limit": ({"nondetect": "mdl"}, "carried at zero"),
    "non-detects carried at half the detection limit": ({"nondetect": "half"}, "carried at zero"),
    "holding-time result kept in the August nickel average": ({"include_invalid": True}, "invalidated the nickel result"),
    "resample left out of the September averages": ({"resample_in_month": False}, "counts in the month's average"),
    "resample left out of the six-month measurements": ({"resample_in_snc": False}, "resamples included"),
    "technical review criterion of 1.4 applied to metals": ({"trc": "1.4"}, "technical review criterion of 1.2"),
    "24 hours counted from the plant manager being told": ({"notice_from_manager": True}, "from the 07:50 event"),
}

EXPECTED = {
    "Jul Cu monthly average": "2.930", "Jul Cu daily maximum": "4.090",
    "Aug Cu monthly average": "4.240", "Aug Cu daily maximum": "4.330",
    "Sep Cu monthly average": "2.053", "Sep Cu daily maximum": "2.960", "Sep Cu valid samples": "3",
    "Aug Ni monthly average": "2.610", "Aug Ni valid samples": "1",
    "Sep Zn monthly average": "1.653", "Sep Zn daily maximum": "3.140",
    "Sep Cd monthly average": "0.004",
    "Sep CN monthly average": "0.325", "Sep CN valid samples": "2",
    "SNC Cu measurements": "13", "SNC Cu above limit": "6", "SNC Cu above TRC": "5",
    "SNC Cu share above TRC pct": "38.5", "SNC Cu TRC threshold": "4.056",
    "SNC Zn above limit": "1", "SNC Ni measurements": "12",
    "September highest daily flow": "47300", "September days above the flow limit": "1",
    "hours from the overflow to the telephone notice": "32.67",
    "resample due after VP-260909 Zn": "10/15/2026",
    "resample due after VP-260721 Cu": "08/29/2026",
    "September Cu average without the resample": "2.320",
}

if __name__ == "__main__":
    rep = Report()
    figures, standings = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures.get(label), golden)
    rep.expect("SNC Cu", standings["SNC Cu"], "Significant noncompliance")
    rep.expect("SNC Zn", standings["SNC Zn"], "Not in significant noncompliance")
    rep.expect("Aug Ni frequency status", standings["Aug Ni frequency status"], "Below frequency")
    rep.expect("overflow telephone notice", standings["overflow telephone notice"], "Late")
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
