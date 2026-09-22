"""verify_golden.py for freight-invoice-audit: re-derive every figure the audit states from inputs/ alone
under the contract as the golden applies it, and list the alternate readings a reviewer could take with
the phrase the golden uses to settle each one.

    .venv/bin/python submissions/07-freight-invoice-audit/verify_golden.py [--print]

The seven inputs: the invoice and bill of lading CSVs, the contract and its rates exhibit, the fuel
bulletin, Northline's billing file transmittal (scale tickets, an inspection certificate, a reclass
notice and a dock readout, read from the particulars of each document) and the shipping desk's email
printout (the services confirmed to Northline in writing, read from the messages themselves).
"""
import csv
import datetime as dt
import os
import re
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
GOLDEN = HERE / "solution" / "northline_audit_sept2026.xlsx"


def r2(x):
    return D(str(x)).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(x):,.2f}"


def pdate(s):
    return dt.datetime.strptime(s, "%m/%d/%Y").date()


# ---------------------------------------------------------------- tariff exhibit
tw = openpyxl.load_workbook(INPUTS / "northline_contract_rates_2026.xlsx", data_only=True)
ws = tw["Lane Rates"]
LANES = {ws.cell(row=r, column=1).value: [D(str(ws.cell(row=r, column=c).value)) for c in range(2, 7)] for r in range(5, 12)}
BREAK_MIN = [int(ws.cell(row=14, column=c).value) for c in range(2, 7)]
ws = tw["Class Factors"]
CLASS_FACTOR = {D(str(ws.cell(row=r, column=1).value)): D(str(ws.cell(row=r, column=2).value)) for r in range(4, 11)}
GROUP_CLASS = {ws.cell(row=r, column=1).value: D(str(ws.cell(row=r, column=2).value)) for r in range(15, 21)}
ws = tw["Accessorials"]
ACC_FLAG = {ws.cell(row=r, column=2).value: D(str(ws.cell(row=r, column=3).value)) for r in range(4, 8)}
ACC_NAME = {ws.cell(row=r, column=1).value: D(str(ws.cell(row=r, column=3).value)) for r in (4, 5, 6, 7, 9, 10)}
DISCOUNT = D(str(ws["C13"].value)); MIN_CHARGE = D(str(ws["C14"].value))
ws = tw["Fuel Surcharge"]
BANDS = [(D(str(ws.cell(row=r, column=1).value)), D(str(ws.cell(row=r, column=3).value))) for r in range(4, 22)]
TOLERANCE = D("0.05")  # section 4.5, "more than five percent"

# ---------------------------------------------------------------- bulletin: DOE price and the printed percentage per Monday
DOE, BULLETIN_PCT = {}, {}
for t in Document(INPUTS / "northline_fuel_bulletin_sept_2026.docx").tables:
    for row in t.rows[1:]:
        monday = dt.datetime.strptime(row.cells[0].text.strip(), "%B %d, %Y").date()
        DOE[monday] = D(row.cells[1].text.strip())
        BULLETIN_PCT[monday] = D(row.cells[2].text.strip().rstrip("%")) / 100

# ---------------------------------------------------------------- billing file: what each listed document is
CERT_TICKET, INSPECTION_CERT, TICKET_WEIGHT = set(), set(), {}
for t in Document(INPUTS / "northline_billing_file_sept_2026.docx").tables:
    for row in t.rows[1:]:
        pro, doc, part = row.cells[0].text.strip(), row.cells[2].text.strip().lower(), row.cells[3].text.lower()
        if doc == "scale ticket" and "certified ticket" in part and "signed by the weighmaster" in part:
            CERT_TICKET.add(pro)
            TICKET_WEIGHT[pro] = int(re.search(r"net ([\d,]+) pounds", part).group(1).replace(",", ""))
        if doc == "inspection certificate" and "signed by" in part and "inspector" in part:
            INSPECTION_CERT.add(pro)

# ---------------------------------------------------------------- correspondence: services the shipper confirmed in writing
CONFIRMED = []  # (pro, service, confirmed date)
paras = [p.text.strip() for p in Document(INPUTS / "shipping_desk_correspondence_sept_2026.docx").paragraphs]
subject_pro, sender, sent = None, None, None
for text in paras:
    m = re.match(r"Subject: Pro (\d{6})", text)
    if m:
        subject_pro = m.group(1); continue
    if text.startswith("From: "):
        sender = text[6:]; continue
    if text.startswith("Sent: "):
        sent = dt.datetime.strptime(text[6:].split(",", 1)[1].strip().rsplit(",", 1)[0], " %B %d, %Y").date() if False else \
            dt.datetime.strptime(re.search(r"(\w+ \d{1,2}, \d{4})", text).group(1), "%B %d, %Y").date()
        continue
    if subject_pro and sender and "Halvorsen" in sender:
        if re.search(r"please add liftgate delivery", text, re.I):
            CONFIRMED.append((subject_pro, "Liftgate delivery", sent))
        if re.search(r"please reconsign", text, re.I):
            CONFIRMED.append((subject_pro, "Reconsignment", sent))

BOLS = {b["PRO_NO"]: b for b in csv.DictReader(open(INPUTS / "bills_of_lading_sept_2026.csv", newline=""))}
INV = list(csv.DictReader(open(INPUTS / "northline_invoices_sept_2026.csv", newline="")))


def fsc_pct(price):
    pct = BANDS[0][1]
    for lo, p in BANDS:
        if price >= lo:
            pct = p
    return pct


def monday(d):
    return d - dt.timedelta(days=d.weekday())


def derive(deficit=True, fsc_on="pickup", fsc_source="table", tickets="apply", certs="apply", confirmations="apply", net_undercharges=False):
    """tickets: apply (certified and over the tolerance holds), all (every certified ticket holds), ignore (no ticket holds),
    billed (every billed weight holds). certs / confirmations: apply or ignore."""
    fig, standing = {}, {}
    seen = set(); rows = []
    for inv in INV:
        pro = inv["PRO_NO"]; b = BOLS[pro]; pickup = pdate(b["PICKUP_DATE"]); delivered = pdate(b["DELIVERED_DATE"])
        dup = pro in seen; seen.add(pro)
        bol_w = int(b["WEIGHT_LB"]); billed_w = int(inv["BILLED_WEIGHT"])
        cert = pro in CERT_TICKET and TICKET_WEIGHT[pro] == billed_w
        over_tol = D(billed_w) > D(bol_w) * (1 + TOLERANCE)
        if tickets == "apply":
            w = billed_w if (cert and over_tol) else bol_w
        elif tickets == "all":
            w = billed_w if cert else bol_w
        elif tickets == "billed":
            w = billed_w
        else:
            w = bol_w
        group_cls = GROUP_CLASS[b["PRODUCT_GROUP"]]; billed_cls = D(inv["BILLED_CLASS"])
        cls = billed_cls if (certs == "apply" and pro in INSPECTION_CERT) else group_cls
        f = CLASS_FACTOR[cls]
        i = max(k for k, m in enumerate(BREAK_MIN) if w >= m); rates = LANES[b["DEST_STATE"]]
        own = r2(D(w) / 100 * rates[i] * f)
        gross = own
        if deficit and i < 4:
            gross = min(own, r2(D(BREAK_MIN[i + 1]) / 100 * rates[i + 1] * f))
        net = max(r2(gross * (1 - DISCOUNT)), MIN_CHARGE)
        fdate = pickup if fsc_on == "pickup" else pdate(inv["INVOICE_DATE"])
        week = monday(fdate)
        pct = fsc_pct(DOE[week]) if fsc_source == "table" else BULLETIN_PCT[week]
        fsc = r2(net * pct)
        acc_bol = sum((ACC_FLAG[k] for k in ACC_FLAG if b[k] == "Y"), D(0))
        acc_conf = D(0)
        if confirmations == "apply":
            acc_conf = sum((ACC_NAME[svc] for p, svc, d in CONFIRMED if p == pro and d < delivered), D(0))
        acc = acc_bol + acc_conf
        expected = D(0) if dup else net + fsc + acc
        billed = D(inv["INVOICE_TOTAL"]); var = billed - expected
        if dup: code = "7.1 duplicate bill"
        elif billed_cls != cls: code = "4.3 class"
        elif billed_w != w: code = "4.5 reweigh"
        elif D(inv["BASE_CHARGE"]) != gross: code = "4.2 deficit weight"
        elif D(inv["FSC_PCT"]) / 100 != pct: code = "5.2 fuel week"
        elif D(inv["ACCESSORIAL_AMT"]) > acc: code = "6.1 accessorial not requested"
        elif D(inv["ACCESSORIAL_AMT"]) < acc: code = "6.1 accessorial omitted"
        else: code = ""
        if certs == "apply" and pro in INSPECTION_CERT and billed_cls != group_cls: settled = "4.3 inspection certificate"
        elif tickets != "ignore" and cert and billed_w != bol_w and billed_w == w: settled = "4.5 certified scale ticket"
        elif acc_conf > 0: settled = "6.1 written confirmation"
        else: settled = ""
        rows.append(dict(inv=inv["INVOICE_NO"], pro=pro, expected=expected, billed=billed, var=var, code=code, settled=settled,
                         net=net, fsc=fsc, acc=acc, pct=pct, gross=gross, w=w, cls=cls))
        standing[inv["INVOICE_NO"]] = "over" if var > 0 else "under" if var < 0 else "correct"
    over = [r for r in rows if r["var"] > 0]; under = [r for r in rows if r["var"] < 0]
    fig["invoices"] = len(rows); fig["billed total"] = money(sum(r["billed"] for r in rows)); fig["contract total"] = money(sum(r["expected"] for r in rows))
    fig["overbilled"] = money(sum(r["var"] for r in over)); fig["underbilled"] = money(-sum(r["var"] for r in under))
    fig["claim total"] = money(sum(r["var"] for r in over) - (sum(-r["var"] for r in under) if net_undercharges else 0))
    fig["claims"] = len(over); fig["undercharges"] = len(under); fig["exceptions"] = sum(1 for r in rows if r["code"]); fig["clean"] = sum(1 for r in rows if not r["code"])
    fig["settled"] = sum(1 for r in rows if r["settled"])
    for c in ("7.1", "4.3", "4.5", "4.2", "5.2", "6.1"):
        fig[f"code {c}"] = sum(1 for r in rows if r["code"].startswith(c))
    fig["net difference"] = money(sum(r["var"] for r in rows))
    fig["rows"] = rows
    standing["claim total"] = fig["claim total"]; standing["settled"] = fig["settled"]
    return fig, standing


VARIANTS = {
    "deficit weight rule not applied": ({"deficit": False}, "lesser of"),
    "fuel surcharge on the invoice-date week": ({"fsc_on": "invoice"}, "pickup date"),
    "bulletin percentage used where it disagrees with the table": ({"fsc_source": "bulletin"}, "the table governs"),
    "every certified ticket held whatever the tolerance": ({"tickets": "all"}, "five percent"),
    "carrier reweigh held without a scale ticket": ({"tickets": "billed"}, "bill of lading weight"),
    "certified scale ticket ignored": ({"tickets": "ignore"}, "certified scale ticket"),
    "inspection certificate ignored": ({"certs": "ignore"}, "inspection certificate"),
    "written confirmation ignored": ({"confirmations": "ignore"}, "confirmed in writing"),
    "undercharges netted against the claim": ({"net_undercharges": True}, "not netted"),
}

if __name__ == "__main__":
    fig, standing = derive()
    if "--print" in sys.argv:
        for k, v in fig.items():
            if k != "rows": print(k, v)
        print("certified tickets", sorted(CERT_TICKET), "inspection certificates", sorted(INSPECTION_CERT), "confirmed", CONFIRMED)
        for r in fig["rows"]:
            if r["code"] or r["var"] or r["settled"]: print(r["inv"], r["pro"], r["code"] or "-", r["settled"] or "-", money(r["billed"]), money(r["expected"]), money(r["var"]))
        sys.exit(0)
    wb = openpyxl.load_workbook(GOLDEN, data_only=True)
    s = wb["Summary"]
    rep = Report()
    summary = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value for r in range(1, 40) if s.cell(row=r, column=1).value}
    rep.expect("invoices", fig["invoices"], summary["Freight bills in the file"])
    rep.expect("billed total", fig["billed total"], money(summary["Total billed by Northline"]))
    rep.expect("contract total", fig["contract total"], money(summary["Total due under the contract"]))
    rep.expect("net difference", fig["net difference"], money(summary["Net difference, billed less due"]))
    rep.expect("overbilled", fig["overbilled"], money(summary["Overcharges claimed"]))
    rep.expect("underbilled", fig["underbilled"], money(summary["Undercharges reported, not netted"]))
    rep.expect("claims", fig["claims"], summary["Freight bills claimed"])
    rep.expect("undercharges", fig["undercharges"], summary["Freight bills underbilled"])
    rep.expect("exceptions", fig["exceptions"], summary["Freight bills with an exception"])
    rep.expect("clean", fig["clean"], summary["Freight bills billed as the contract provides"])
    rep.expect("settled", fig["settled"], summary["Billed items accepted on a document or a written confirmation"])
    for c, label in (("7.1", "Duplicate freight bills (7.1)"), ("4.3", "Class not the product group class (4.3)"), ("4.5", "Reweigh that does not hold (4.5)"),
                     ("4.2", "Deficit weight rule not applied (4.2)"), ("5.2", "Fuel surcharge on the wrong week (5.2)"), ("6.1", "Accessorial not requested or omitted (6.1)")):
        rep.expect(f"code {c}", fig[f"code {c}"], summary[label])
    # audit rows: every expected total, code and settlement across both panels
    a = wb["Audit"]
    hdr = {a.cell(row=3, column=c).value: c for c in range(1, a.max_column + 1) if a.cell(row=3, column=c).value}
    for k, r in enumerate(fig["rows"]):
        panel = "" if k < 19 else "_2"; row = 4 + (k if k < 19 else k - 19)
        rep.expect(f"{r['inv']} key", r["inv"], a.cell(row=row, column=hdr["INVOICE_NO" + panel]).value)
        rep.expect(f"{r['inv']} expected", money(r["expected"]), money(a.cell(row=row, column=hdr["CONTRACT_TOTAL" + panel]).value))
        rep.expect(f"{r['inv']} variance", money(r["var"]), money(a.cell(row=row, column=hdr["VARIANCE" + panel]).value))
        rep.expect(f"{r['inv']} code", r["code"], a.cell(row=row, column=hdr["EXCEPTION" + panel]).value or "")
        rep.expect(f"{r['inv']} settled", r["settled"], a.cell(row=row, column=hdr["SETTLED_BY" + panel]).value or "")
        rep.expect(f"{r['inv']} contract weight", r["w"], a.cell(row=row, column=hdr["CONTRACT_WEIGHT" + panel]).value)
        rep.expect(f"{r['inv']} contract class", r["cls"], D(str(a.cell(row=row, column=hdr["CONTRACT_CLASS" + panel]).value)))
    cl = wb["Claim Schedule"]
    rep.expect("claim schedule total", fig["claim total"], money(cl.cell(row=4 + fig["claims"], column=5).value))
    rep.expect("claim schedule rows", fig["claims"], sum(1 for r in range(4, 4 + fig["claims"]) if cl.cell(row=r, column=2).value))
    un = wb["Undercharges"]
    rep.expect("undercharge total", fig["underbilled"], money(un.cell(row=4 + fig["undercharges"], column=5).value))
    note = "\n".join(str(c.value) for row in wb["Note to Ingrid"].iter_rows() for c in row if c.value)
    for phrase in (f"${fig['overbilled']}", f"${fig['underbilled']}", f"${fig['billed total']}", f"${fig['contract total']}", f"{fig['claims']} freight bills"):
        rep.expect(f"note states {phrase}", phrase in note, True)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
