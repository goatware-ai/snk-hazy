"""verify_golden.py for carton-bid-evaluation: re-derive every figure the evaluation states from inputs/ alone
under ITB 2026-17 as amended and policy PP-3 as the golden applies them (six bids, the bidders' letters, the
quality hold register and the current drawing revisions), and list the alternate readings a reviewer could take
with the phrase the golden uses to settle each one.

    .venv/bin/python submissions/10-carton-bid-evaluation/verify_golden.py
"""
import os
import re
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
GOLDEN = HERE / "solution" / "bid_evaluation_itb_2026_17.xlsx"


def r2(x):
    return D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(x):,.2f}"


# ---------------------------------------------------------------- inputs
tw = openpyxl.load_workbook(INPUTS / "bid_tabulation_itb_2026_17.xlsx", data_only=True)
ws = tw["Bids"]
KEYS = [ws.cell(row=4, column=c).value.replace("_UNIT", "") for c in range(3, 15, 2)]
PRICES = {k: [D(str(ws.cell(row=r, column=3 + 2 * i).value)) for r in range(5, 10)] for i, k in enumerate(KEYS)}
WRITTEN = {k: [D(str(ws.cell(row=r, column=4 + 2 * i).value)) for r in range(5, 10)] for i, k in enumerate(KEYS)}
TOOLING = {k: D(str(ws.cell(row=10, column=4 + 2 * i).value)) for i, k in enumerate(KEYS)}
ws = tw["Bidder Terms"]
TERMS = {}
for i, k in enumerate(KEYS):
    r = 4 + i
    name = ws.cell(row=r, column=1).value
    assert name.split()[0].upper() == k.split("_")[0], (name, k)
    TERMS[k] = dict(name=name, state=ws.cell(row=r, column=3).value, fob=ws.cell(row=r, column=4).value,
                    pay=ws.cell(row=r, column=5).value, bond=ws.cell(row=r, column=6).value,
                    addenda=ws.cell(row=r, column=8).value, samples=ws.cell(row=r, column=9).value)
ws = tw["Freight Schedule"]
FREIGHT = {ws.cell(row=r, column=1).value: D(str(ws.cell(row=r, column=2).value)) for r in range(4, 11)}
ws = tw["Item Weights"]
LB_CURRENT, LB_SUPERSEDED = {}, {}
for r in range(4, 10):
    item, status, lb = ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value, D(str(ws.cell(row=r, column=5).value))
    (LB_CURRENT if status.startswith("Current") else LB_SUPERSEDED)[item] = lb
assert set(LB_CURRENT) == {1, 2, 3, 4, 5} and LB_SUPERSEDED == {5: D("1.4")}

itb = "\n".join(p.text for p in Document(INPUTS / "itb_2026_17_cartons.docx").paragraphs)
add = "\n".join(p.text for p in Document(INPUTS / "itb_2026_17_addendum_2.docx").paragraphs)
pol = "\n".join(p.text for p in Document(INPUTS / "purchasing_policy_pp3_excerpt.docx").paragraphs)
QTY = [60000, 84000, 48000, 12000, 30000]
QTY_SUPERSEDED_3 = 48000
QTY[2] = int(re.search(r"reduced from 48,000 to ([\d,]+)", add).group(1).replace(",", ""))
DISC_MIN_DAYS = 20
APPROVAL = 100000
assert "twenty days" in itb and "one hundred thousand" in pol
assert "conditions its price on truckload releases takes an exception to the quantity term" in add
assert "open quality hold" in pol and "passed over" in pol and "revision B replaces revision A" in add

# the letters: each bidder's conditions are read from the bid forms document
forms = Document(INPUTS / "bid_forms_itb_2026_17.docx")
LETTER = {}
current = None
for p in forms.paragraphs:
    m = re.match(r"Letter enclosed with the bid, (.+)", p.text)
    if m:
        current = m.group(1)
        LETTER[current] = ""
        continue
    if p.text.startswith("Bid form, "):
        current = None
    if current:
        LETTER[current] += p.text + "\n"
assert set(LETTER) == {t["name"] for t in TERMS.values()}


def conditions(name):
    t = LETTER[name].lower()
    return dict(price_hold=("six months" in t and "index" in t),
                truckload=("full truckload" in t and "stop charge" in t),
                alternate=("alternate" in t))


# the hold register: open holds against a bidder, matched on the registered name's leading words
hw = openpyxl.load_workbook(INPUTS / "supplier_quality_holds_2026.xlsx", data_only=True)["Holds"]
HOLDS = []
for r in range(4, hw.max_row + 1):
    if hw.cell(row=r, column=1).value and str(hw.cell(row=r, column=1).value).startswith("QH-"):
        HOLDS.append(dict(no=hw.cell(row=r, column=1).value, supplier=hw.cell(row=r, column=2).value, status=hw.cell(row=r, column=6).value))


def open_holds(name):
    return [h["no"] for h in HOLDS if h["supplier"].startswith(name) and h["status"] == "Open"]


def parse_terms(t):
    m = re.match(r"(\d+)% (\d+)", t)
    return (D(m.group(1)) / 100, int(m.group(2))) if m else (D(0), 0)


def derive(unit_governs=True, discount_min_days=DISC_MIN_DAYS, freight_on="origin", price_hold_material=True,
           truckload_material=True, addendum_material=True, hold_bars=True, item5_rev="current"):
    fig, standing = {}, {}
    ev = {}
    lb = dict(LB_CURRENT)
    if item5_rev == "A":
        lb[5] = LB_SUPERSEDED[5]
    lb_total = sum(lb[i + 1] * QTY[i] for i in range(5))
    cwt = lb_total / 100
    for k in KEYS:
        t = TERMS[k]
        c = conditions(t["name"])
        ext = [r2(PRICES[k][i] * QTY[i]) if unit_governs else WRITTEN[k][i] for i in range(5)]
        sub = sum(ext)
        freight = r2(cwt * FREIGHT[t["state"]]) if (t["fob"].startswith("Origin") and freight_on == "origin") else D(0)
        rate, days = parse_terms(t["pay"])
        credit = r2(sub * rate) if days >= discount_min_days else D(0)
        total = sub + TOOLING[k] + freight - credit
        reasons = []
        if t["bond"] != "Yes":
            reasons.append("no bid bond")
        if "2" not in t["addenda"] and addendum_material:
            reasons.append("addendum 2 not acknowledged; bid on the superseded item 3 quantity and item 5 grade")
        if t["samples"] != "Yes":
            reasons.append("samples not received")
        if WRITTEN[k][2] == r2(PRICES[k][2] * QTY_SUPERSEDED_3) and "2" in t["addenda"] and addendum_material:
            reasons.append("item 3 priced at the superseded quantity")
        if c["price_hold"] and price_hold_material:
            reasons.append("exception to the twelve month price hold, a material term")
        if c["truckload"] and truckload_material:
            reasons.append("price conditioned on truckload releases, an exception to the quantity term")
        responsive = not reasons
        ev[k] = dict(ext=ext, sub=sub, freight=freight, credit=credit, total=total, responsive=responsive,
                     reason="; ".join(reasons), written=sum(WRITTEN[k]) + TOOLING[k], holds=open_holds(t["name"]))
        fig[f"{k} subtotal"] = money(sub)
        fig[f"{k} freight"] = money(freight)
        fig[f"{k} credit"] = money(credit)
        fig[f"{k} total"] = money(total)
        fig[f"{k} responsive"] = "Responsive" if responsive else "Non-responsive"
        fig[f"{k} written total"] = money(ev[k]["written"])
        standing[f"{k} responsive"] = fig[f"{k} responsive"]
    resp = sorted([k for k in KEYS if ev[k]["responsive"]], key=lambda k: ev[k]["total"])
    award, passed, nxt = None, None, None
    for n, k in enumerate(resp):
        if hold_bars and ev[k]["holds"]:
            passed = k
            continue
        award = k
        nxt = resp[n + 1] if n + 1 < len(resp) else None
        break
    fig["award"] = TERMS[award]["name"]
    fig["award total"] = money(ev[award]["total"])
    fig["runner up"] = TERMS[nxt]["name"] if nxt else "None"
    fig["margin"] = money(ev[nxt]["total"] - ev[award]["total"]) if nxt else "None"
    fig["passed over"] = TERMS[passed]["name"] if passed else "None"
    fig["passed total"] = money(ev[passed]["total"]) if passed else "None"
    fig["passed below award"] = money(ev[award]["total"] - ev[passed]["total"]) if passed else "None"
    fig["passed hold"] = ", ".join(ev[passed]["holds"]) if passed else "None"
    fig["approver"] = "Vice president of operations" if ev[award]["total"] > APPROVAL else "Director of supply chain"
    low_written = min(KEYS, key=lambda k: ev[k]["written"])
    fig["lowest as written"] = TERMS[low_written]["name"]
    fig["responsive count"] = len(resp)
    fig["bids received"] = len(KEYS)
    fig["freight lb"] = money(lb_total)
    fig["hundredweight"] = money(cwt)
    for k in KEYS:
        for i in range(5):
            if ev[k]["ext"][i] != WRITTEN[k][i]:
                fig[f"corrected {k} item {i + 1}"] = f"{money(WRITTEN[k][i])} to {money(ev[k]['ext'][i])}"
    standing["award"] = fig["award"]
    standing["approver"] = fig["approver"]
    standing["passed over"] = fig["passed over"]
    return fig, standing


VARIANTS = {
    "bidder's written extensions carried": ({"unit_governs": False}, "unit price governs"),
    "discount credited at any period": ({"discount_min_days": 0}, "twenty days or longer"),
    "freight not added for FOB origin": ({"freight_on": "none"}, "FOB origin"),
    "price hold condition waived": ({"price_hold_material": False}, "material term"),
    "truckload release condition waived": ({"truckload_material": False}, "quantity term"),
    "unacknowledged addendum waived": ({"addendum_material": False}, "non-responsive"),
    "open quality hold ignored": ({"hold_bars": False}, "not responsible"),
    "item 5 weighed at the superseded revision A": ({"item5_rev": "A"}, "revision B"),
}

if __name__ == "__main__":
    fig, standing = derive()
    if "--print" in sys.argv:
        for k, v in fig.items():
            print(k, v)
        sys.exit(0)
    wb = openpyxl.load_workbook(GOLDEN, data_only=True)
    rep = Report()
    s = wb["Recommendation"]
    rows = {s.cell(row=r, column=1).value: [s.cell(row=r, column=c).value for c in range(2, 8)] for r in range(1, 40) if s.cell(row=r, column=1).value}
    hdr = rows["Bidder"]
    for j, name in enumerate(hdr):
        key = KEYS[j]
        assert TERMS[key]["name"] == name, (key, name)
        rep.expect(f"{key} responsive", fig[f"{key} responsive"], rows["Responsiveness"][j])
        rep.expect(f"{key} written total", fig[f"{key} written total"], money(rows["Total bid as written"][j]))
        rep.expect(f"{key} subtotal", fig[f"{key} subtotal"], money(rows["Corrected extensions"][j]))
        if fig[f"{key} responsive"] == "Responsive":
            rep.expect(f"{key} freight", fig[f"{key} freight"], money(rows["Freight added, FOB origin"][j]))
            rep.expect(f"{key} credit", fig[f"{key} credit"], money(rows["Payment discount credit"][j]))
            rep.expect(f"{key} total", fig[f"{key} total"], money(rows["Total evaluated price"][j]))
    rep.expect("award", fig["award"], rows["Recommended award"][0])
    rep.expect("award total", fig["award total"], money(rows["Recommended award, total evaluated price"][0]))
    rep.expect("runner up", fig["runner up"], rows["Next responsive bidder in line"][0])
    rep.expect("margin", fig["margin"], money(rows["Margin over the next responsive bid"][0]))
    rep.expect("passed over", fig["passed over"], rows["Responsive bid passed over, open quality hold"][0])
    rep.expect("passed total", fig["passed total"], money(rows["Passed-over bid, total evaluated price"][0]))
    rep.expect("passed below award", fig["passed below award"], money(rows["Passed-over bid below the recommended award by"][0]))
    rep.expect("approver", fig["approver"], rows["Approver required by PP-3.5.3"][0])
    rep.expect("lowest as written", fig["lowest as written"], rows["Lowest bid as written"][0])
    rep.expect("responsive count", fig["responsive count"], rows["Responsive bids"][0])
    rep.expect("bids received", fig["bids received"], len([h for h in hdr if h]))
    passed_col = hdr.index(fig["passed over"])
    rep.expect("passed over responsibility", "Not responsible, open quality hold, PP-3.7.1", rows["Responsibility of the bidder in line, PP-3.7.1"][passed_col])
    fr = wb["Freight"]
    rep.expect("freight lb", fig["freight lb"], money(fr["H9"].value))
    rep.expect("hundredweight", fig["hundredweight"], money(fr["H10"].value))
    note = "\n".join(str(c.value) for row in wb["Note to Torsten"].iter_rows() for c in row if c.value)
    for phrase in (f"${fig['award total']}", f"${fig['margin']}", f"${fig['passed below award']}", f"${fig['passed total']}",
                   fig["passed hold"], "45,112.00", "35,112.00", f"${fig['RIDGECREST freight']}", f"${fig['RIDGECREST credit']}",
                   "1,917 hundredweight", "revision B"):
        rep.expect(f"note states {phrase}", phrase in note, True)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
