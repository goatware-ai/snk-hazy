"""verify_golden.py for title-exam-cedarbrook-lot12: re-derive every figure and every standing the
examination report states from inputs/ alone, under the conventions the report states, and list the
alternate readings a reviewer could take with the phrase the report uses to settle each one.

    .venv/bin/python drafts/02-title-exam-cedarbrook-lot12/verify_golden.py
"""
import csv
import datetime as dt
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
CLOSING = dt.date(2026, 11, 6)


def c(x):
    return D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(x):,.2f}"


def pdate(s):
    return dt.datetime.strptime(s, "%m/%d/%Y").date()


INDEX = list(csv.DictReader(open(INPUTS / "recorder_index_cedarbrook_lot12.csv", newline="")))
wb = openpyxl.load_workbook(INPUTS / "tax_and_lien_search_hlt262214.xlsx", data_only=True)
TAX = [r for r in wb["Tax Duplicate"].iter_rows(values_only=True) if isinstance(r[0], int)]
JL = [r for r in wb["Judgment Liens"].iter_rows(values_only=True) if isinstance(r[0], str) and r[0][2:4] == "JL"]
NAMES = [r for r in wb["Names"].iter_rows(values_only=True) if isinstance(r[0], str) and r[1] and r[0] != "NAME"]


def derive(penalty="0.10", year_days=365, proration_through_closing=True, surname_attach=False,
           sale_releases=False, prorate_assessment=False, ignore_variance=False):
    pen = D(penalty)
    # chain of title: every warranty deed in order, the land each describes
    deeds = [r for r in INDEX if r["DOC_TYPE"] == "WARRANTY DEED"]
    gap = [r for r in deeds if "N 1/2 Lot 13" not in r["LEGAL"]]
    gap_deed = gap[0] if gap else None
    north_half_owner = gap_deed["GRANTOR"] if gap_deed else deeds[-1]["GRANTEE"]
    # mortgages with no satisfaction referencing them
    released = {r["REF_INSTRUMENT"] for r in INDEX if r["DOC_TYPE"].startswith("SATISFACTION") or r["DOC_TYPE"].startswith("RELEASE")}
    mortgages = [r for r in INDEX if r["DOC_TYPE"] == "MORTGAGE"]
    open_mtgs = [r["INSTRUMENT_NO"] for r in mortgages if r["INSTRUMENT_NO"] not in released]
    if sale_releases:
        # a reader who assumes the 2014 sale paid the 2011 loan
        open_mtgs = [m for m in open_mtgs if not m.startswith("2011")]
    liens = [r for r in INDEX if r["DOC_TYPE"] == "MECHANICS LIEN"]
    open_liens = [r["INSTRUMENT_NO"] for r in liens if r["INSTRUMENT_NO"] not in released]
    # judgment liens: attach only where the debtor's name matches an owner holding on or after the filing date
    periods = {}
    for n in NAMES:
        to = CLOSING if n[3] == "current" else pdate(n[3])
        periods[n[0]] = to
    attach = {}
    for row in JL:
        debtor, filed = row[3], pdate(row[1])
        hit = periods.get(debtor)
        if surname_attach and hit is None:
            surname = debtor.split()[0]
            hit = max((v for k, v in periods.items() if k.split()[0] == surname), default=None)
        attach[row[0]] = "attaches" if hit is not None and hit >= filed else "does not attach"
    # the seller's judgment payoff
    seller = next(row for row in JL if attach[row[0]] == "attaches")
    amount, costs = D(str(seller[5])), D(str(seller[6]))
    rate = D(seller[7].split("%")[0]) / 100
    days = (CLOSING - pdate(seller[1])).days
    interest = c(amount * rate * days / year_days)
    payoff = amount + interest + costs
    # delinquent installments with penalty
    delinquent = [r for r in TAX if r[7] == "Delinquent"]
    tax_due = sum((D(str(r[3])) for r in delinquent), D(0))
    ass_due = sum((D(str(r[4])) for r in delinquent), D(0))
    delinq_total = tax_due + c(tax_due * pen) + ass_due + c(ass_due * pen)
    # proration on the most recent full-year bill
    last_year = max(r[0] for r in TAX)
    annual = sum((D(str(r[3])) for r in TAX if r[0] == last_year), D(0))
    if prorate_assessment:
        annual += sum((D(str(r[4])) for r in TAX if r[0] == last_year), D(0))
    pro_days = (CLOSING - dt.date(CLOSING.year, 1, 1)).days + (1 if proration_through_closing else 0)
    proration = c(annual * pro_days / 365)
    # dower and name variance from the abstracts are facts of the instruments, read here from the index names
    grantee_2014 = next(r["GRANTEE"] for r in deeds if r["INSTRUMENT_NO"] == "201409190112")
    grantor_2018 = next(r["GRANTOR"] for r in deeds if r["INSTRUMENT_NO"] == "201805040071")
    variance = "affidavit of identity required" if (grantee_2014 != grantor_2018 and not ignore_variance) else "none"
    figures = {
        "judgment interest": money(interest),
        "judgment payoff": money(payoff),
        "interest days": str(days),
        "delinquent total": money(delinq_total),
        "tax penalty": money(c(tax_due * pen)),
        "assessment penalty": money(c(ass_due * pen)),
        "annual tax estimate": money(annual),
        "proration days": str(pro_days),
        "seller tax credit": money(proration),
        "open mortgage 2011 amount": money(next(D(r["AMOUNT"]) for r in mortgages if r["INSTRUMENT_NO"] == "201103140062")),
    }
    standings = {
        "north half of Lot 13 record owner": north_half_owner,
        "deed describing Lot 12 only": gap_deed["INSTRUMENT_NO"] if gap_deed else "none",
        "open mortgages of record": ",".join(open_mtgs),
        "open mechanics liens": ",".join(open_liens) or "none",
        "20JL00417": attach["20JL00417"],
        "25JL01133": attach["25JL01133"],
        "judgment payoff": money(payoff),
        "delinquent total": money(delinq_total),
        "seller tax credit": money(proration),
        "Kessler name variance": variance,
    }
    return figures, standings


VARIANTS = {
    "penalty at 5% instead of 10%": ({"penalty": "0.05"}, "10% penalty"),
    "judgment interest on a 360-day year": ({"year_days": 360}, "365-day year"),
    "proration to the day before closing": ({"proration_through_closing": False}, "through the day of closing"),
    "judgment attached on a surname match": ({"surname_attach": True}, "before the filing date"),
    "2011 mortgage assumed paid at the 2014 sale": ({"sale_releases": True}, "open of record"),
    "special assessment prorated with the tax": ({"prorate_assessment": True}, "The special assessment is not prorated"),
    "Peggy and Margaret Ann read as one name": ({"ignore_variance": True}, "one and the same person"),
}

EXPECTED = {
    "judgment interest": "957.19",
    "judgment payoff": "10,556.19",
    "interest days": "464",
    "delinquent total": "2,319.48",
    "tax penalty": "179.61",
    "assessment penalty": "31.25",
    "annual tax estimate": "3,592.24",
    "proration days": "310",
    "seller tax credit": "3,050.94",
    "open mortgage 2011 amount": "88,000.00",
}

if __name__ == "__main__":
    rep = Report()
    figures, standings = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures.get(label), golden)
    rep.expect("north half of Lot 13 record owner", standings["north half of Lot 13 record owner"], "Pruitt Daniel R")
    rep.expect("open mortgages of record", standings["open mortgages of record"], "201103140062,201805040072")
    rep.expect("20JL00417", standings["20JL00417"], "does not attach")
    rep.expect("25JL01133", standings["25JL01133"], "attaches")
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
