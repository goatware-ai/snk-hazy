"""verify_golden.py for concrete-acceptance-review: re-derive every figure the summary states from inputs/ alone
under the specification and the laboratory procedure as the golden applies them, and list the alternate readings
a reviewer could take with the phrase the golden uses to settle each one.

    .venv/bin/python drafts/09-concrete-acceptance-review/verify_golden.py
"""
import csv
import math
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
GOLDEN = HERE / "solution" / "acceptance_summary_2417_aug_sep.xlsx"
FC = {"M-4000": 4000, "M-6000": 6000}
AREA = {6: D("28.27"), 4: D("12.57")}
CYL = list(csv.DictReader(open(INPUTS / "cylinder_breaks_2417.csv", newline="")))
PL = list(csv.DictReader(open(INPUTS / "placement_log_2417.csv", newline="")))
MIX = {p["PLACEMENT_ID"]: p["MIX_ID"] for p in PL}
memo = "\n".join(p.text for p in Document(INPUTS / "lab_quality_memo_2417.docx").paragraphs)
CURING = [w for w in memo.replace(",", " ").split() if w.startswith("S-") and "curing box" in memo.split(w, 1)[1][:200]]
CURING_SET = "S-14"
assert "S-14" in memo and "91 F" in memo
spec = "\n".join(p.text for p in Document(INPUTS / "spec_033000_excerpt.docx").paragraphs)
assert "20 hours" in spec and "500 psi" in spec and "10 percent" in spec


def r10(x):
    return int((D(x) / 10).quantize(D(1), rounding=ROUND_HALF_UP) * 10)


def derive(area_mode="procedure", age_tol_hours=20, count_rule=True, avg_of_rounded=True, curing_counts=True):
    fig, standing = {}, {}
    sets = {}
    for c in CYL:
        d = int(c["DIAMETER_IN"])
        area = AREA[d] if area_mode == "procedure" else D(str(round(math.pi * d * d / 4, 6)))
        strength = r10(D(c["MAX_LOAD_LBF"]) / area)
        s = sets.setdefault(c["SET_ID"], dict(placement=c["PLACEMENT_ID"], size=d, cast=c["CAST_DATE"], age=D(c["AGE_DAYS"]), strengths=[]))
        s["strengths"].append(strength)
    order = sorted(sets, key=lambda k: (sets[k]["cast"][6:] + sets[k]["cast"][:5], k))
    for sid in order:
        s = sets[sid]
        vals = s["strengths"]; s["test"] = r10(D(sum(vals)) / len(vals)) if avg_of_rounded else r10(D(sum(vals)) / len(vals))
        s["required"] = 3 if s["size"] == 4 else 2
        s["range_pct"] = (D(max(vals) - min(vals)) / D(s["test"])).quantize(D("0.0001"))
        s["range_limit"] = D("0.106") if s["size"] == 4 else D("0.066")
        s["range_note"] = s["range_pct"] > s["range_limit"]
        s["age_ok"] = abs(s["age"] - 28) <= D(age_tol_hours) / 24
        s["count_ok"] = len(vals) >= s["required"] if count_rule else True
        s["curing"] = sid == CURING_SET
        s["used"] = s["age_ok"] and s["count_ok"] and (curing_counts or not s["curing"])
        s["mix"] = MIX[s["placement"]]
        fig[f"{sid} test"] = s["test"]; fig[f"{sid} used"] = "Yes" if s["used"] else "No"
    for mix, fc in FC.items():
        used = [sid for sid in order if sets[sid]["mix"] == mix and sets[sid]["used"]]
        allow = 500 if fc <= 5000 else fc // 10
        a_fail, b_fail, movs = [], [], {}
        for i, sid in enumerate(used):
            t = sets[sid]["test"]
            if t < fc - allow: b_fail.append(sid)
            if i >= 2:
                m = (D(sets[used[i - 2]]["test"] + sets[used[i - 1]]["test"] + t) / 3).quantize(D("0.1"), rounding=ROUND_HALF_UP)
                movs[sid] = m
                if m < fc: a_fail.append(sid)
        fig[f"{mix} tests used"] = len(used); fig[f"{mix} tests not used"] = sum(1 for sid in order if sets[sid]["mix"] == mix and not sets[sid]["used"])
        fig[f"{mix} a failures"] = len(a_fail); fig[f"{mix} b failures"] = len(b_fail)
        fig[f"{mix} lowest test"] = min(sets[sid]["test"] for sid in used)
        fig[f"{mix} b failing placements"] = ", ".join(sets[sid]["placement"] for sid in b_fail)
        fig[f"{mix} status"] = "Meets 3.11.1" if not a_fail and not b_fail else "Does not meet 3.11.1"
        fig[f"{mix} core average floor"] = int(fc * 0.85); fig[f"{mix} core single floor"] = int(fc * 0.75)
        for sid, m in movs.items(): fig[f"{sid} moving average"] = str(m)
        standing[f"{mix} status"] = fig[f"{mix} status"]; standing[f"{mix} a failures"] = len(a_fail); standing[f"{mix} b failures"] = len(b_fail)
        standing[f"{mix} tests used"] = len(used)
    # sampling
    short = []
    for p in PL:
        req = math.ceil(int(p["CUBIC_YARDS"]) / 150); cast = sum(1 for s in sets.values() if s["placement"] == p["PLACEMENT_ID"])
        fig[f"{p['PLACEMENT_ID']} required"] = req; fig[f"{p['PLACEMENT_ID']} cast"] = cast
        if cast < req: short.append(p["PLACEMENT_ID"])
    fig["sampling short placements"] = ", ".join(short); fig["sampling short count"] = len(short)
    fig["sets not used"] = sum(1 for s in sets.values() if not s["used"]); fig["range noted"] = sum(1 for s in sets.values() if s["range_note"])
    fig["curing noted"] = sum(1 for s in sets.values() if s["curing"])
    fresh = []
    for p in PL:
        lim = {"M-4000": (3, 5, 4.5, 7.5, 90), "M-6000": (5, 7, 2.0, 4.0, 90)}[p["MIX_ID"]]
        ok = lim[0] <= float(p["SLUMP_IN"]) <= lim[1] and lim[2] <= float(p["AIR_PCT"]) <= lim[3] and float(p["CONCRETE_TEMP_F"]) <= lim[4]
        if not ok: fresh.append(p["PLACEMENT_ID"])
    fig["fresh outside count"] = len(fresh); fig["fresh outside placements"] = ", ".join(fresh)
    standing["sampling short count"] = len(short); standing["sets not used"] = fig["sets not used"]
    fig["_sets"] = sets; fig["_order"] = order
    return fig, standing


VARIANTS = {
    "area from pi rather than the procedure's stated areas": ({"area_mode": "pi"}, "28.27 and 12.57 square inches"),
    "age tolerance of 24 hours": ({"age_tol_hours": 24}, "20 hours"),
    "two 4 by 8 cylinders accepted as a test": ({"count_rule": False}, "three 4 by 8"),
    "curing excursion set left out": ({"curing_counts": False}, "counted for acceptance"),
}

if __name__ == "__main__":
    fig, standing = derive()
    if "--print" in sys.argv:
        for k, v in fig.items():
            if not k.startswith("_"): print(k, v)
        sys.exit(0)
    wb = openpyxl.load_workbook(GOLDEN, data_only=True)
    rep = Report()
    s = wb["Summary"]
    summ = {}
    for r in range(1, 60):
        lab = s.cell(row=r, column=1).value
        if lab: summ[lab] = (s.cell(row=r, column=2).value, s.cell(row=r, column=3).value)
    for j, mix in enumerate(("M-4000", "M-6000")):
        rep.expect(f"{mix} tests used", fig[f"{mix} tests used"], summ["Strength tests used for acceptance"][j])
        rep.expect(f"{mix} tests not used", fig[f"{mix} tests not used"], summ["Sets cast but not used"][j])
        rep.expect(f"{mix} a failures", fig[f"{mix} a failures"], summ["Moving averages of three below f'c, criterion (a)"][j])
        rep.expect(f"{mix} b failures", fig[f"{mix} b failures"], summ["Tests below f'c by more than the allowance, criterion (b)"][j])
        rep.expect(f"{mix} lowest test", fig[f"{mix} lowest test"], summ["Lowest strength test (psi)"][j])
        rep.expect(f"{mix} status", fig[f"{mix} status"], summ["Acceptance under 3.11.1"][j])
        rep.expect(f"{mix} b failing placements", fig[f"{mix} b failing placements"], summ["Placement of each test failing criterion (b)"][j])
        rep.expect(f"{mix} core average floor", fig[f"{mix} core average floor"], summ["Core average floor, 85 percent of f'c (psi)"][j])
    rep.expect("sampling short count", fig["sampling short count"], summ["Placements sampled short of 3.10.1"][0])
    rep.expect("sets not used", fig["sets not used"], summ["Sets cast but not used, both mixtures"][0])
    rep.expect("range noted", fig["range noted"], summ["Tests with a within-test range note"][0])
    rep.expect("curing noted", fig["curing noted"], summ["Sets with a curing excursion note"][0])
    rep.expect("fresh outside count", fig["fresh outside count"], summ["Placements outside a fresh concrete limit"][0])
    t = wb["Tests"]
    hdr = {t.cell(row=3, column=c).value: c for c in range(1, t.max_column + 1) if t.cell(row=3, column=c).value}
    for r in range(4, 4 + len(fig["_order"])):
        sid = t.cell(row=r, column=1).value
        rep.expect(f"{sid} test", fig[f"{sid} test"], t.cell(row=r, column=hdr["STRENGTH_TEST_PSI"]).value)
        rep.expect(f"{sid} used", fig[f"{sid} used"], t.cell(row=r, column=hdr["USED_FOR_ACCEPTANCE"]).value)
    for name in ("Acceptance M-4000", "Acceptance M-6000"):
        a = wb[name]
        for r in range(4, 16):
            sid = a.cell(row=r, column=1).value
            if sid and a.cell(row=r, column=4).value not in (None, ""):
                rep.expect(f"{sid} moving average", fig[f"{sid} moving average"], str(D(str(a.cell(row=r, column=4).value)).quantize(D("0.1"))))
    sp = wb["Sampling"]
    for r in range(4, 4 + len(PL)):
        pid = sp.cell(row=r, column=1).value
        rep.expect(f"{pid} required", fig[f"{pid} required"], sp.cell(row=r, column=5).value)
        rep.expect(f"{pid} cast", fig[f"{pid} cast"], sp.cell(row=r, column=6).value)
    note = "\n".join(str(c.value) for row in wb["Note to Gunnar"].iter_rows() for c in row if c.value)
    for phrase in ("3,480 psi", "5,370 psi", "P-05", "P-10", "3,400", "5,100", "S-13", "S-12"):
        rep.expect(f"note states {phrase}", phrase in note, True)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
