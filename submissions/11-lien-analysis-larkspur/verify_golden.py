"""verify_golden.py for lien-analysis-larkspur: re-derive every figure the analysis states from inputs/ alone
under Practice Note PN-14 as the golden applies it (the recorder's copies of the sixteen claims with the statements
of account three suppliers attached and the proofs of service two claimants attached, the owner's project file with
its notice, waiver, payment, certified mail and returned mail records, the direct contractor's check ledger, and the
close-out file with its court notice), and list the alternate readings a solver could take with the phrase the golden
uses to settle each one.

    .venv/bin/python submissions/11-lien-analysis-larkspur/verify_golden.py [--print [--variants]]
"""
import csv
import math
import os
import re
import sys
import datetime as dt
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(
    (str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file()), "/Users/aladdin/projects/snk/hazy")
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = Path(os.environ.get("LIEN_INPUTS") or HERE / "inputs")
GOLDEN = Path(os.environ.get("LIEN_GOLDEN") or HERE / "solution" / "lien_analysis_larkspur_crossing.xlsx")
MONTHS = {m: i for i, m in enumerate(["January", "February", "March", "April", "May", "June", "July", "August",
                                       "September", "October", "November", "December"], start=1)}


def r2(x):
    return D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(str(x)):,.2f}"


def longdate(s):
    m = re.search(r"([A-Z][a-z]+) (\d{1,2}), (\d{4})", s)
    return dt.date(int(m.group(3)), MONTHS[m.group(1)], int(m.group(2)))


def mdy(s):
    m, d, y = s.split("/")
    return dt.date(int(y), int(m), int(d))


def fmt(d):
    return d.strftime("%m/%d/%Y")


def dollars(s):
    return D(s.replace("$", "").replace(",", ""))


def doc_text(path):
    d = Document(path)
    return "\n".join(p.text for p in d.paragraphs), d


# ---------------------------------------------------------------- practice note: the rules and constants
note, _ = doc_text(INPUTS / "practice_note_pn14_mechanics_liens.docx")
for phrase in ("within 15 days after completion", "within 10 days after it is recorded", "30 days after the owner records a notice of completion",
               "60 days after the owner records a notice of completion", "ineffective to shorten the time",
               "the 90 days after completion is that claimant's period", "not later than 20 days after the claimant first furnishes",
               "within 90 days after the claim is recorded", "125 percent of the amount of the claim of lien as recorded",
               "rounded up to the next whole dollar", "the sum of the bonds so written",
               "at least 10 days before the petition", "Saturday, a Sunday, or a holiday", "closed for the whole of the day is a holiday",
               "the day the notice is given is not counted", "effective when signed, whether or not the claimant has been paid",
               "only once the bank has paid it", "A payment the claimant receives after it records reduces what it can enforce", "finance charges",
               "a change order request that was not approved adds nothing to it", "saves only a disputed claim for extras that the form states in a dollar amount",
               "A document in any other form does not release the lien", "retention included", "direct contractual relationship with the owner",
               "only on a search run after the last day", "unenforceable as a matter of law", "at the address shown on the building permit",
               "accepts a preliminary notice mailed to the owner at the address the permit shows",
               "as adjusted by the change orders the owner has approved", "within 45 days after completion", "withhold 150 percent",
               "less 150 percent of any amount in good faith dispute, and never less than zero", "2 percent per month on that amount"):
    assert phrase in note, phrase
NOC_DAYS, COPY_DAYS, SUB_DAYS, DC_DAYS, NO_NOC_DAYS, PRELIM_DAYS, ACTION_DAYS, BOND_PCT = 15, 10, 30, 60, 90, 20, 90, D(125)
RET_DAYS, DISPUTE_PCT, PENALTY_PCT = 45, D(150), D(2)

# ---------------------------------------------------------------- close-out file
corr, cdoc2 = doc_text(INPUTS / "completion_and_correspondence_larkspur.docx")
COMPLETION = longdate(re.search(r"was completed on ([A-Z][a-z]+ \d+, \d{4})", corr).group(0))
NOC_RECORDED = longdate(re.search(r"Recorded ([A-Z][a-z]+ \d+, \d{4}), Placer County Recorder, Document No\. 2026-0081562", corr).group(0))
INDEX_DATE = longdate(re.search(r"civil case index search, run ([A-Z][a-z]+ \d+, \d{4})", corr).group(0))
assert "no civil action found naming Piedmont Ridge Properties" in corr and "no notice of pendency of action" in corr
assert "no notice of credit" in corr and "no release of any claim" in corr
NO_ACTION = True
RETENTION = dollars(re.search(r"will stand at (\$[\d,]+\.\d\d) once application 8 is certified", corr).group(1))
ADJUSTED = dollars(re.search(r"six approved change orders, (\$[\d,]+\.\d\d) against the original", corr).group(1))
CONTRACT_SIGNED = dollars(re.search(r"against the original (\$[\d,]+\.\d\d)", corr).group(1))
HVAC = dollars(re.search(r"estimated the correction at (\$[\d,]+\.\d\d)", corr).group(1))
assert "disputes that the shortfall is its responsibility" in corr
m = re.search(r"up to (\d+) percent of the value in the bank's appraisal, which came in at (\$[\d,]+\.\d\d)", corr)
LINE_PCT, APPRAISAL = D(m.group(1)), dollars(m.group(2))
SURETY_LINE = r2(APPRAISAL * LINE_PCT / 100)
PERMIT_ADDRESS = re.search(r"old office, (1440 Eureka Road, Suite 100), as the owner's address on that permit", corr).group(1)
BANK_DATE = "November 6" in corr
assert "how much of Dunmore-Kettle's retention we have to release now" in corr
recorder_rows = [[c.text for c in r.cells] for r in cdoc2.tables[0].rows[1:]]
notice_e = corr[corr.index("E. Public notice of office closures"):]
HOLIDAYS = {longdate(m.group(1)) for m in re.finditer(r"closed on (?:Monday|Tuesday|Wednesday|Thursday|Friday), ([A-Z][a-z]+ \d{1,2}, \d{4}), Labor Day", notice_e)}
CLOSURES = {longdate(m.group(1)) for m in re.finditer(r"closed to the public for the whole of (?:Monday|Tuesday|Wednesday|Thursday|Friday), ([A-Z][a-z]+ \d{1,2}, \d{4})", notice_e)}
assert HOLIDAYS == {dt.date(2026, 9, 7)} and CLOSURES == {dt.date(2026, 9, 25)}, (HOLIDAYS, CLOSURES)
search_para = re.search(r"Placer County Superior Court, civil case index search.*?described\.", corr, re.S).group(0)
assert "closed" not in search_para and "Labor Day" not in corr.split("E. Public notice")[0].split("C. Correspondence")[1]

# ---------------------------------------------------------------- the claims
ctext, cdoc = doc_text(INPUTS / "claims_of_lien_larkspur_crossing.docx")
CLAIMS = {}
for m in re.finditer(r"Claim (\d+)\. (.+?), recorded ([A-Z][a-z]+ \d+, \d{4}), Document No\. (\S+)\n(.*?)(?=\nClaim \d+\. |\Z)", ctext, re.S):
    n, name, rec, doc, body = m.groups()
    amount = dollars(re.search(r"after deducting all just credits and offsets, is (\$[\d,]+\.\d\d)", body).group(1))
    first = longdate(re.search(r"first furnished work on ([A-Z][a-z]+ \d+, \d{4})", body).group(0))
    last = longdate(re.search(r"last furnished work on ([A-Z][a-z]+ \d+, \d{4})", body).group(0))
    pm = re.search(r"preliminary notice was given on ([A-Z][a-z]+ \d+, \d{4})", body)
    item5 = re.search(r"\n5\. (.*?)\n6\.", body, re.S).group(1)
    item6 = re.search(r"6\. General description of the work furnished: (.*?)\n7\.", body, re.S).group(1)
    service = re.search(r"I served a copy of this claim of mechanics lien on .*?postage prepaid, addressed to (.*?)\. I declare", body, re.S).group(1)
    CLAIMS[name] = dict(n=int(n), recorded=longdate(rec), doc=doc, amount=amount, first=first, last=last,
                        recital=longdate(pm.group(0)) if pm else None, item5=item5, item6=item6,
                        direct="at the request of Piedmont Ridge Properties, LLC, the owner" in item5,
                        served_owner_office=service.startswith("the owner at Piedmont Ridge Properties") and "3100 Douglas Boulevard" in service,
                        served_permit=service.startswith("the owner at Piedmont Ridge Properties") and PERMIT_ADDRESS in service and "building permit" in service)
assert len(CLAIMS) == 16
assert sorted(r[0] for r in recorder_rows if r[2] == "Claim of mechanics lien") == sorted(c["doc"] for c in CLAIMS.values())
POS, STATEMENTS = {}, {}
for tbl in cdoc.tables:
    first_cell = tbl.rows[1].cells[0].text
    hdr = [c.text for c in tbl.rows[0].cells]
    if hdr[0] == "Party":
        parties = [(r.cells[0].text, r.cells[1].text) for r in tbl.rows[1:]]
        owner = next((a for p_, a in parties if "Piedmont" in p_ or "owner" in p_.lower()), None)
        key = "Granite Bay Sheet Metal, Inc." if first_cell.startswith("Piedmont") else "Norcal Rebar & Mesh, LLC"
        POS[key] = dict(owner_address=owner)
    else:
        assert hdr == ["Invoice", "Invoice date", "Delivery ticket", "Delivered", "Materials", "Amount", "Paid", "Balance"]
        key = {"S": "Hedrick & Sons Building Supply", "F": "Sierra Pipe & Supply, Inc.", "C": "Capitol Wire & Lighting, Inc."}[first_cell[0]]
        lines = []
        for r in tbl.rows[1:-1]:
            inv, invdate, ticket, delivered, mat, amt, paid, bal = [c.text for c in r.cells]
            lines.append(dict(inv=inv, invdate=mdy(invdate), ticket=ticket, date=mdy(delivered) if delivered else None,
                              amount=dollars(amt), balance=dollars(bal), charge="inance charge" in mat))
        assert sum(h["balance"] for h in lines) == CLAIMS[key]["amount"], key
        STATEMENTS[key] = lines
        CLAIMS[key]["first_work"] = min(h["date"] for h in lines if h["date"])
for name, c in CLAIMS.items():
    c.setdefault("first_work", c["first"])
assert POS["Norcal Rebar & Mesh, LLC"]["owner_address"] is None and PERMIT_ADDRESS in POS["Granite Bay Sheet Metal, Inc."]["owner_address"]
NOTICE_EXHIBIT_DATE = {"Granite Bay Sheet Metal, Inc.": longdate(re.search(r"I, Rosa Almeida, declare that on ([A-Z][a-z]+ \d+, \d{4}) I served the attached", ctext).group(1)),
                       "Norcal Rebar & Mesh, LLC": longdate(re.search(r"I, Danielle Ruiz, declare that on ([A-Z][a-z]+ \d+, \d{4}) I served the attached", ctext).group(1))}

# ---------------------------------------------------------------- owner's project file
wb = openpyxl.load_workbook(INPUTS / "owner_project_file_larkspur.xlsx", data_only=True)
PRELIM = {}
for r in wb["Prelim Log"].iter_rows(min_row=5, values_only=True):
    if r[0]:
        PRELIM[r[0]] = dict(on_notice=mdy(r[3]), mailed=mdy(r[4]), received=mdy(r[5]), first=mdy(r[6]))
assert "Granite Bay Sheet Metal, Inc." not in PRELIM and "Norcal Rebar & Mesh, LLC" not in PRELIM
SUBVAL, SUBCO = {}, {}
for r in wb["Subcontractors"].iter_rows(min_row=4, values_only=True):
    if r[0]:
        SUBVAL[r[0]] = D(str(r[3]))
        SUBCO[r[0]] = D(str(r[4] or 0))
WAIVERS = []
for r in wb["Waiver Log"].iter_rows(min_row=5, values_only=True):
    for block in (r[0:9], r[9:18]):
        if block[0]:
            WAIVERS.append(dict(claimant=block[0], form=block[2], through=mdy(block[3]) if block[3] else None,
                                amount=D(str(block[4])), exceptions=block[5], check=block[6]))
assert len(WAIVERS) == 40
ws = wb["Owner Payments"]
head = ws["A1"].value
CO = dollars(re.search(r"approved change orders 1 to 6 (\$[\d,]+\.\d\d)", head).group(1))
assert dollars(re.search(r"contract sum (\$[\d,]+\.\d\d), approved", head).group(1)) == CONTRACT_SIGNED
assert CONTRACT_SIGNED + CO == ADJUSTED
paid_dk = sum(D(str(r[5])) for r in ws.iter_rows(min_row=4, max_row=12, values_only=True) if r[5] is not None)
certified = sum(D(str(r[3])) for r in ws.iter_rows(min_row=4, max_row=12, values_only=True) if r[3] is not None)
retained = sum(D(str(r[4])) for r in ws.iter_rows(min_row=4, max_row=12, values_only=True) if r[4] is not None)
app8 = next(r for r in ws.iter_rows(min_row=4, max_row=12, values_only=True) if r[0] == 8)
APP8_NET = D(str(app8[3])) - D(str(app8[4]))
assert certified == ADJUSTED and retained == RETENTION, (certified, retained)
HELD = ADJUSTED - paid_dk
HELD_SIGNED = CONTRACT_SIGNED - paid_dk
assert HELD == APP8_NET + RETENTION
DIRECT_VENDORS = {}
for r in ws.iter_rows(min_row=18, max_row=24, values_only=True):
    if r[0]:
        DIRECT_VENDORS[r[0]] = dict(balance=D(str(r[6])), waiver=r[7])
        m = re.match(r"(Unconditional|Conditional) Waiver and Release on (Progress|Final) Payment through (\d\d/\d\d/\d{4})", r[7] or "")
        if m:
            WAIVERS.append(dict(claimant=r[0], form=f"{m.group(1)} Waiver and Release on {m.group(2)} Payment",
                                through=mdy(m.group(3)), amount=D(str(r[4])), exceptions="None", check=None))
MAIL, RETURNED = [], {}
cm_rows = list(wb["Certified Mail"].iter_rows(min_row=5, values_only=True))
for r in cm_rows:
    if r[0] and "notice of completion" in (r[4] or "").lower():
        MAIL.append(dict(mailed=mdy(r[0]), addressee=r[2], article=r[1]))
    elif r[0] and r[0] != "DATE_RETURNED" and r[2] and "nclaimed" in str(r[2]):
        RETURNED[r[1]] = mdy(r[0])
assert len(RETURNED) == 1

# ---------------------------------------------------------------- the direct contractor's ledger
CHECKS, STOPPED = {}, set()
with open(INPUTS / "dk_subcontractor_payments.csv", encoding="utf-8-sig") as fh:
    _raw = list(csv.reader(fh))
_w = 8
assert _raw[0][_w:] == [h + "_2" for h in _raw[0][:_w]], "the ledger is two side-by-side panels with _2 suffixes"
LEDGER = [dict(zip(_raw[0][:_w], r[:_w])) for r in _raw[1:]] + [dict(zip(_raw[0][:_w], r[_w:])) for r in _raw[1:] if any(r[_w:])]
assert len(LEDGER) == 40
if True:
    for row in LEDGER:
        no = int(row["CHECK_NO"])
        if row["TYPE"] == "CHK":
            CHECKS[no] = row
        elif row["TYPE"] == "STP":
            STOPPED.add(no)


def check_state(no):
    if no in STOPPED:
        return "Stopped"
    return "Paid" if CHECKS[no]["BANK_CLEARED"] else "Not paid"


assert check_state(4460) == "Stopped" and check_state(4461) == "Paid" and check_state(4479) == "Not paid" and check_state(4448) == "Paid" and check_state(4488) == "Paid"


def exception_amount(w):
    m = re.search(r"extras: \$([\d,]+\.\d\d)", w["exceptions"] or "")
    return dollars(m.group(1)) if m else D(0)


# ---------------------------------------------------------------- the method
def extend(d, holidays=None):
    """Code of Civil Procedure sections 12a and 12b: a last day on a Saturday, a Sunday, a holiday or a whole-day office closure runs to the next day that is none of them."""
    holidays = (HOLIDAYS | CLOSURES) if holidays is None else holidays
    while d.weekday() >= 5 or d in holidays:
        d += dt.timedelta(days=1)
    return d


WORDS = {"bldg": "building", "&": "and"}
DROP = {"inc", "llc", "co", "corp", "the"}


def tokens(name):
    out = []
    for w in re.sub(r"[.,]", " ", name.lower()).replace("&", " and ").replace("-", " ").split():
        w = WORDS.get(w, w)
        if w not in DROP:
            out.append(w)
    return out


def contains(seq, sub):
    return any(seq[i:i + len(sub)] == sub for i in range(len(seq) - len(sub) + 1))


def mailed_copy(name, match="full", unclaimed_not_given=False):
    """The date a copy of the notice of completion was deposited for this person, or None."""
    want = tokens(name)
    for row in MAIL:
        have = tokens(row["addressee"])
        hit = have[0] == want[0] if match == "surname" else have == want[:len(have)]
        if hit:
            if unclaimed_not_given and row["article"] in RETURNED:
                return None
            return row["mailed"]
    return None


def paid_later(name, amount, joint_ok=True):
    """A check to the claimant (alone or jointly) for this amount that the bank paid, issued after completion."""
    for no, row in CHECKS.items():
        if no in STOPPED or not row["BANK_CLEARED"]:
            continue
        payee = tokens(row["PAYEE"])
        hit = payee == tokens(name) or (joint_ok and contains(payee, tokens(name)))
        if hit and D(row["AMOUNT"]) == amount and mdy(row["ENTRY_DATE"]) > COMPLETION:
            return no
    return None


def components(name):
    """Each claim built up from its own statement: (label, date furnished through, amount, kind, invoice date)."""
    c = CLAIMS[name]
    out = []
    if name in STATEMENTS:
        for h in STATEMENTS[name]:
            if h["balance"] > 0:
                out.append((f"{h['inv']} {h['ticket']}", h["date"] or h["invdate"], h["balance"], "charges" if h["charge"] else "materials", h["invdate"]))
        return out
    if name in ("Norcal Rebar & Mesh, LLC", "Foothill Ready Mix, Inc.", "Granite Bay Sheet Metal, Inc."):
        return [("materials", c["last"], c["amount"], "materials", c["last"])]
    if name == "Tallac Concrete, Inc.":
        ret = r2(SUBVAL[name] * D("0.05"))
        return [("retention", c["last"], ret, "retention", c["last"]), ("May billing balance", c["last"], c["amount"] - ret, "progress", c["last"])]
    t = c["item6"]
    for m in re.finditer(r"(?:progress payment|final billing|final progress payment) for the period ending ([A-Z][a-z]+ \d+, \d{4}) in the sum of (\$[\d,]+\.\d\d)", t):
        out.append((m.group(0)[:40], longdate(m.group(1)), dollars(m.group(2)), "progress", longdate(m.group(1))))
    m = re.search(r"final invoice for work performed .*? through ([A-Z][a-z]+ \d+, \d{4}) in the sum of (\$[\d,]+\.\d\d)", t)
    if m:
        out.append(("final invoice", longdate(m.group(1)), dollars(m.group(2)), "progress", longdate(m.group(1))))
    m = re.search(r"retention withheld in the sum of (\$[\d,]+\.\d\d)", t)
    out.append(("retention", c["last"], dollars(m.group(1)), "retention", c["last"]))
    m = re.search(r"invoice (\S+) dated .*? through ([A-Z][a-z]+ \d+, \d{4}) in the sum of (\$[\d,]+\.\d\d)", t)
    if m:
        out.append((m.group(1), longdate(m.group(2)), dollars(m.group(3)), "extra", longdate(m.group(2))))
    m = re.search(r"change order request (\d+) dated .*? through ([A-Z][a-z]+ \d+, \d{4}) in the sum of (\$[\d,]+\.\d\d)", t)
    if m:
        out.append((f"change order request {m.group(1)}", longdate(m.group(2)), dollars(m.group(3)), "extra", longdate(m.group(2))))
    assert sum(x[2] for x in out) == c["amount"], (name, out, c["amount"])
    return out


def derive(weekend_ext=True, closure_is_holiday=True, later_payment=True, joint_ok=True, charges_out=True, copy_rule=True, copy_ext=True, copy_grace=0,
           mail_match="full", unclaimed_not_given=False, direct_recognised=True, notice_date="mailed", first_from="statement", furnished="delivered",
           prelim_days=PRELIM_DAYS, noc_inclusive=False, cond_on_signing=False, issued_is_paid=False, uncond_needs_payment=False,
           final_saves_retention=False, final_exception=True, other_form_releases=False, price_rule=True, service_checked=True,
           permit_address_ok=True, permit_notice_ok=True, dc_notice_suffices=False, expiry="search", bond_pct=BOND_PCT, bond_round="each",
           deadline_from="noc", held_basis="adjusted", dispute_pct=DISPUTE_PCT, claims_withheld=True):
    fig, standing = {}, {}
    hol = HOLIDAYS | (CLOSURES if closure_is_holiday else set())
    ext = (lambda d: extend(d, hol)) if weekend_ext else (lambda d: d)
    fig["days completion to noc"] = (NOC_RECORDED - COMPLETION).days
    noc_ok = fig["days completion to noc"] + (1 if noc_inclusive else 0) <= NOC_DAYS
    base = NOC_RECORDED if deadline_from == "noc" else COMPLETION
    last_sub, last_dc = ext(base + dt.timedelta(days=SUB_DAYS)), ext(base + dt.timedelta(days=DC_DAYS))
    last_none = ext(COMPLETION + dt.timedelta(days=NO_NOC_DAYS))
    copy_by = NOC_RECORDED + dt.timedelta(days=COPY_DAYS)
    if copy_ext:
        copy_by = extend(copy_by, hol)
    copy_by += dt.timedelta(days=copy_grace)
    held = HELD if held_basis == "adjusted" else HELD_SIGNED
    fig["last day to give copy"] = fmt(copy_by)
    fig["last day others"] = fmt(last_sub)
    fig["last day direct contractor with notice"] = fmt(last_dc)
    fig["last day without notice"] = fmt(last_none)
    fig["retention due"] = fmt(extend(COMPLETION + dt.timedelta(days=RET_DAYS), hol))
    fig["withhold for hvac"] = money(r2(HVAC * dispute_pct / 100))
    fig["held on the direct contract"] = money(held)
    fig["surety line"] = money(SURETY_LINE)

    def last_day(name, is_dc, entitled):
        if not noc_ok:
            return last_none
        if copy_rule and entitled:
            sent = mailed_copy(name, mail_match, unclaimed_not_given)
            if sent is None or sent > copy_by:
                return last_none
        return last_dc if is_dc else last_sub

    dk_last = last_day("Dunmore-Kettle Builders, Inc.", True, True)
    fig["last day Dunmore-Kettle"] = fmt(dk_last)
    total_rec = total_enf = total_bond = total_bond_exact = release_amt = chain_enf = D(0)
    n_bond = n_release = 0
    for name, c in CLAIMS.items():
        is_dc = c["direct"] and direct_recognised
        # ---- lien rights
        in_log = name in PRELIM
        if is_dc:
            rights, coverage, given = True, c["first_work"], None
        else:
            if in_log:
                given = {"mailed": PRELIM[name]["mailed"], "on_notice": PRELIM[name]["on_notice"],
                         "received": PRELIM[name]["received"], "recital": c["recital"] or PRELIM[name]["mailed"]}[notice_date]
                rights = True
            elif name in POS:
                given = NOTICE_EXHIBIT_DATE[name]
                owner_addr = POS[name]["owner_address"]
                rights = bool(owner_addr) and (PERMIT_ADDRESS not in owner_addr or permit_notice_ok)
                rights = rights or bool(dc_notice_suffices and given)
            else:
                given = c["recital"]
                rights = bool(dc_notice_suffices and given)
            first = {"statement": c["first_work"], "claim": c["first"], "log": PRELIM[name]["first"] if in_log else c["first"]}[first_from]
            if given is None:
                coverage = first
            else:
                coverage = first if (given - first).days <= prelim_days else given - dt.timedelta(days=PRELIM_DAYS)
        entitled = is_dc or rights
        # ---- recording deadline
        ld = last_day(name, is_dc, entitled)
        timely = c["recorded"] <= ld
        # ---- service
        served = c["served_owner_office"] or (c["served_permit"] and permit_address_ok) or not service_checked
        # ---- expiry
        due = ext(c["recorded"] + dt.timedelta(days=ACTION_DAYS))
        if expiry == "search":
            expired = NO_ACTION and INDEX_DATE > due
        elif expiry == "due date alone":
            expired = NO_ACTION and dt.date(2026, 10, 20) > due
        else:
            expired = False

        # ---- waivers
        def effective(w):
            title = w["form"]
            statutory = title.startswith("Unconditional Waiver and Release") or title.startswith("Conditional Waiver and Release")
            if not statutory:
                return other_form_releases
            if title.startswith("Unconditional"):
                return not uncond_needs_payment
            m = re.search(r"check (\d{4})", w["check"] or "")
            state = check_state(int(m.group(1))) if m else "Not paid"
            if state == "Paid" or cond_on_signing:
                return True
            return state == "Not paid" and issued_is_paid

        comps = components(name)
        unpaid_sum = sum(amt for label, date, amt, kind, inv in comps
                         if kind != "charges" and not (later_payment and kind in ("progress", "materials") and paid_later(name, amt, joint_ok)))

        def released(date, kind):
            for w in WAIVERS:
                if w["claimant"] != name or not effective(w):
                    continue
                if "Final Payment" in w["form"]:
                    if kind == "retention" and final_saves_retention:
                        continue
                    if final_exception and exception_amount(w) >= unpaid_sum:
                        continue
                    return True
                if kind in ("retention", "extra"):
                    continue
                if w["through"] is not None and w["through"] >= date:
                    return True
            return False

        enf = D(0)
        for label, date, amt, kind, invdate in comps:
            basis = date if furnished == "delivered" else invdate
            if kind == "charges":
                if not charges_out:
                    enf += amt
                continue
            if kind == "extra" and price_rule and SUBCO.get(name, D(0)) < amt:
                continue  # a change order request never approved is no part of the price agreed (section 8430(a))
            if kind in ("progress", "materials") and later_payment and paid_later(name, amt, joint_ok):
                continue
            if basis >= coverage and not released(date, kind):
                enf += amt
        if not rights:
            st = "No lien rights"
        elif not timely:
            st = "Untimely"
        elif not served:
            st = "Not served on the owner"
        elif expired:
            st = "Expired"
        elif enf == 0:
            st = "Waived in full"
        else:
            st = "Enforceable"
        enforceable = enf if st == "Enforceable" else D(0)
        exact = c["amount"] * bond_pct / 100 if st == "Enforceable" else D(0)
        bond = D(math.ceil(exact)) if (bond_round == "each" and st == "Enforceable") else r2(exact)
        total_rec += c["amount"]
        total_enf += enforceable
        total_bond += bond
        total_bond_exact += exact
        if st == "Enforceable":
            n_bond += 1
            if not c["direct"]:
                chain_enf += enforceable
        else:
            n_release += 1
            release_amt += c["amount"]
        fig[f"{name} standing"] = st
        fig[f"{name} enforceable"] = money(enforceable)
        fig[f"{name} bond"] = money(bond)
        fig[f"{name} timely"] = "Yes" if timely else "No"
        fig[f"{name} rights"] = "Yes" if rights else "No"
        fig[f"{name} served"] = "Yes" if served else "No"
        fig[f"{name} expired"] = "Yes" if expired else "No"
        fig[f"{name} last day"] = fmt(ld)
        fig[f"{name} action due"] = fmt(due)
        fig[f"{name} coverage from"] = fmt(coverage)
        standing[f"{name} standing"] = st
        standing[f"{name} enforceable"] = money(enforceable)
        standing[f"{name} bond"] = money(bond)
    if bond_round == "total":
        total_bond = D(math.ceil(total_bond_exact))
    elif bond_round == "cent":
        total_bond = r2(total_bond_exact)
    withhold = r2(HVAC * dispute_pct / 100)
    release = max(D(0), r2(held - (chain_enf if claims_withheld else D(0)) - withhold))
    fig["claims as recorded"] = money(total_rec)
    fig["enforceable total"] = money(total_enf)
    fig["bonds total"] = money(total_bond)
    fig["release demands amount"] = money(release_amt)
    fig["claims to bond"] = n_bond
    fig["claims to demand released"] = n_release
    fig["enforceable under the direct contract"] = money(chain_enf)
    fig["bonds over the line"] = money(total_bond - SURETY_LINE)
    fig["bonds fit the line"] = "Yes" if total_bond <= SURETY_LINE else "No"
    fig["held over the enforceable claims"] = money(held - chain_enf)
    fig["held covers"] = "Yes" if chain_enf <= held else "No"
    fig["noc effective"] = "Yes" if noc_ok else "No"
    fig["retention to release"] = money(release)
    fig["penalty per month"] = money(r2(release * PENALTY_PCT / 100))
    for k in ("last day others", "last day Dunmore-Kettle", "bonds total", "enforceable under the direct contract", "bonds fit the line", "held covers",
              "claims to bond", "noc effective", "held on the direct contract", "held over the enforceable claims", "retention to release", "penalty per month"):
        standing[k] = fig[k]
    return fig, standing


VARIANTS = {
    "no extension of a last day at all": ({"weekend_ext": False}, "closed to the public for the whole of that day"),
    "the September 25 closure not treated as a holiday": ({"closure_is_holiday": False}, "closed to the public for the whole of that day"),
    "fifteen days to record the notice counted with the day of completion": ({"noc_inclusive": True}, "the fifteenth day after the completion"),
    "replacement check not found on the ledger": ({"later_payment": False}, "check 4474"),
    "joint check not read as payment to the supplier": ({"joint_ok": False}, "joint check 4488"),
    "final form's stated exception ignored": ({"final_exception": False}, "disputed extras it states in a dollar amount"),
    "finance charges left in the enforceable amount": ({"charges_out": False}, "finance charges"),
    "copy of the notice of completion never checked": ({"copy_rule": False}, "was never given a copy of the notice of completion"),
    "tenth day for the copy not extended": ({"copy_ext": False}, "Labor Day"),
    "a copy mailed the day after the last day treated as given in time": ({"copy_grace": 1}, "one day after the last day to give it"),
    "a copy returned unclaimed treated as not given": ({"unclaimed_not_given": True}, "unclaimed or not"),
    "mail log matched on the surname alone": ({"mail_match": "surname"}, "Calder Bros. Masonry is a different company"),
    "contract with the owner not recognised": ({"direct_recognised": False}, "contracted with the owner itself"),
    "notice date read from the date typed on the notice": ({"notice_date": "on_notice"}, "deposited on May 29"),
    "notice date read from the date received": ({"notice_date": "received"}, "deposited on May 29"),
    "first furnishing read from the preliminary notice form": ({"first_from": "log"}, "first delivery was April 6"),
    "first furnishing read from the recital in the claim": ({"first_from": "claim"}, "first delivery of January 26 on its own statement of account"),
    "the twenty-first day treated as within the 20 days": ({"prelim_days": 21}, "the twenty-first day after its first delivery"),
    "deliveries dated by invoice": ({"furnished": "invoice"}, "delivered on May 8"),
    "conditional waiver treated as effective when signed": ({"cond_on_signing": True}, "has not paid at the bank"),
    "issued check treated as payment": ({"issued_is_paid": True}, "has not paid at the bank"),
    "unconditional waiver treated as ineffective without payment": ({"uncond_needs_payment": True}, "effective when signed"),
    "final waiver read as saving retention": ({"final_saves_retention": True}, "retention included"),
    "release on the contractor's own form given effect": ({"other_form_releases": True}, "not one of the four statutory forms"),
    "unapproved change order request treated as part of the price agreed": ({"price_rule": False}, "the request was never approved"),
    "service on the owner not checked": ({"service_checked": False}, "was served on Dunmore-Kettle and not on the owner"),
    "service at the building permit address rejected": ({"permit_address_ok": False}, "address shown on the building permit"),
    "preliminary notice to the owner at the permit address rejected": ({"permit_notice_ok": False}, "given to the owner although the owner's log never recorded it"),
    "notice to the direct contractor alone accepted": ({"dc_notice_suffices": True}, "never given to the owner"),
    "expiry judged by the due date without a later search": ({"expiry": "due date alone"}, "before its last day"),
    "ninety day expiry not checked": ({"expiry": "none"}, "has expired"),
    "bond measured at 100 percent": ({"bond_pct": D(100)}, "125 percent"),
    "bonds rounded once on the total": ({"bond_round": "total"}, "rounded once"),
    "bonds computed to the cent": ({"bond_round": "cent"}, "Computed to the cent"),
    "thirty days counted from completion": ({"deadline_from": "completion"}, "thirtieth day after the notice of completion was recorded"),
    "money held measured against the contract sum as signed": ({"held_basis": "signed"}, "adjusted by the six approved change orders"),
    "dispute hold at 100 percent of the estimate": ({"dispute_pct": D(100)}, "A hold at 100 percent of the estimate"),
    "enforceable claims of record not withheld from what is due": ({"claims_withheld": False}, "enforceable claims of record under the contract"),
}

if __name__ == "__main__":
    fig, standing = derive()
    if "--print" in sys.argv:
        for k, v in fig.items():
            print(k, "|", v)
        if "--variants" in sys.argv:
            for name, (kw, conv) in VARIANTS.items():
                alt = derive(**kw)[1]
                print("==", name, {k: (standing[k], v) for k, v in alt.items() if v != standing[k]})
        sys.exit(0)
    gwb = openpyxl.load_workbook(GOLDEN, data_only=True)
    rep = Report()
    s = gwb[gwb.sheetnames[0]]
    rows = {str(s.cell(row=r, column=1).value): r for r in range(1, s.max_row + 1) if s.cell(row=r, column=1).value}
    kv = lambda label: s.cell(row=rows[label], column=2).value  # noqa: E731
    rep.expect("days completion to noc", fig["days completion to noc"], kv("Days from completion to recording"))
    rep.expect("noc effective", fig["noc effective"], "Yes" if str(kv("Notice of completion effective")).startswith("Yes") else "No")
    rep.expect("last day to give copy", fig["last day to give copy"], kv("Last day to give a copy of the notice of completion"))
    rep.expect("last day others", fig["last day others"], kv("Last day to record, claimant given its copy in time"))
    rep.expect("last day without notice", fig["last day without notice"], kv("Last day to record, person not given a copy in time"))
    rep.expect("last day Dunmore-Kettle", fig["last day Dunmore-Kettle"], kv("Last day for Dunmore-Kettle Builders to record"))
    hdr = rows["CLAIMANT"]
    heads = {s.cell(row=hdr, column=c).value: c for c in range(1, s.max_column + 1) if s.cell(row=hdr, column=c).value}
    for r in range(hdr + 1, hdr + 1 + len(CLAIMS)):
        name = s.cell(row=r, column=1).value
        g = lambda h: s.cell(row=r, column=heads[h]).value  # noqa: E731
        rep.expect(f"{name} standing", fig[f"{name} standing"], g("STANDING"))
        rep.expect(f"{name} rights", fig[f"{name} rights"], g("LIEN_RIGHTS"))
        rep.expect(f"{name} timely", fig[f"{name} timely"], g("TIMELY"))
        rep.expect(f"{name} served", fig[f"{name} served"], g("SERVED_ON_OWNER"))
        rep.expect(f"{name} expired", fig[f"{name} expired"], "Yes" if g("EXPIRED") == "Yes" else "No")
        rep.expect(f"{name} enforceable", fig[f"{name} enforceable"], money(g("ENFORCEABLE_AMOUNT")))
        rep.expect(f"{name} bond", fig[f"{name} bond"], money(g("RELEASE_BOND")))
        rep.expect(f"{name} recorded amount", money(CLAIMS[name]["amount"]), money(g("AMOUNT_AS_RECORDED")))
        rep.expect(f"{name} recorded", fmt(CLAIMS[name]["recorded"]), g("RECORDED"))
    rep.expect("claims as recorded", fig["claims as recorded"], money(kv("Amount of the sixteen claims as recorded")))
    rep.expect("enforceable total", fig["enforceable total"], money(kv("Amount the claimants could enforce against the property")))
    rep.expect("bonds total", fig["bonds total"], money(kv("Release bonds required, total penal sum at 125 percent")))
    rep.expect("release demands amount", fig["release demands amount"], money(kv("Recorded amount of the claims to demand released")))
    rep.expect("claims to bond", fig["claims to bond"], kv("Claims to bond or settle"))
    rep.expect("claims to demand released", fig["claims to demand released"], kv("Claims to demand released"))
    rep.expect("surety line", fig["surety line"], money(kv("Line the surety will write for release bonds")))
    rep.expect("bonds over the line", fig["bonds over the line"], money(kv("Bonds required over the surety's line")))
    rep.expect("bonds fit the line", fig["bonds fit the line"], "Yes" if str(kv("Do the bonds fit inside the surety's line")).startswith("Yes") else "No")
    rep.expect("enforceable under the direct contract", fig["enforceable under the direct contract"],
               money(kv("Enforceable by claimants under the Dunmore-Kettle contract")))
    rep.expect("held on the direct contract", fig["held on the direct contract"], money(kv("Held by the owner on the Dunmore-Kettle contract")))
    rep.expect("held over the enforceable claims", fig["held over the enforceable claims"], money(kv("Money held over the enforceable claims")))
    rep.expect("held covers", fig["held covers"], "Yes" if str(kv("Does what the owner holds cover those claims")).startswith("Yes") else "No")
    rep.expect("retention due", fig["retention due"], kv("Date the retention fell due"))
    rep.expect("retention to release", fig["retention to release"], money(kv("Retention Piedmont Ridge must release to Dunmore-Kettle now")))
    rep.expect("penalty per month", fig["penalty per month"], money(kv("Penalty for each month all of it is held past the due date")))
    # working tabs
    p = gwb["Notices"]
    for r in range(4, 4 + len(CLAIMS)):
        name = p.cell(row=r, column=1).value
        rep.expect(f"{name} coverage from", fig[f"{name} coverage from"], p.cell(row=r, column=12).value)
    dl = gwb["Deadlines"]
    for r in range(4, 4 + len(CLAIMS)):
        name = dl.cell(row=r, column=1).value
        rep.expect(f"{name} last day", fig[f"{name} last day"], dl.cell(row=r, column=4).value)
        rep.expect(f"{name} action due", fig[f"{name} action due"], dl.cell(row=r, column=8).value)
    am = gwb["Amounts"]
    parts = {}
    for r in range(4, am.max_row + 1):
        name = am.cell(row=r, column=1).value
        if name in CLAIMS:
            parts[name] = parts.get(name, D(0)) + D(str(am.cell(row=r, column=6).value))
    for name, c in CLAIMS.items():
        rep.expect(f"{name} components sum to the claim", money(c["amount"]), money(parts.get(name, 0)))
    pr = gwb["Parameters"]
    prm = {pr.cell(row=r, column=1).value: pr.cell(row=r, column=2).value for r in range(4, 70) if pr.cell(row=r, column=1).value}
    rep.expect("retention due (parameters)", fig["retention due"], prm["Retention due to the direct contractor"])
    rep.expect("withhold for hvac", fig["withhold for hvac"], money(prm["Retention the owner may withhold for the rooftop unit correction"]))
    rep.expect("retention withheld", money(RETENTION), money(prm["Retention withheld from the direct contractor"]))
    rep.expect("application 8 net", money(APP8_NET), money(prm["Application 8 unpaid, net of its retention"]))
    rep.expect("paid to the direct contractor", money(paid_dk), money(prm["Paid to the direct contractor to date"]))
    rep.expect("adjusted contract sum", money(ADJUSTED), money(prm["Adjusted contract sum"]))
    rep.expect("change orders approved", money(CO), money(prm["Change orders 1 to 6 approved by the owner"]))
    rep.expect("appraisal", money(APPRAISAL), money(prm["Value in the bank's appraisal"]))
    notetab = gwb[gwb.sheetnames[-1]]
    ntext = "\n".join(str(c.value) for row in notetab.iter_rows() for c in row if c.value)
    for phrase in (f"${fig['bonds total']}", f"${fig['enforceable under the direct contract']}", f"${fig['held on the direct contract']}",
                   f"${fig['held over the enforceable claims']}", f"${money(HELD_SIGNED)}", f"${fig['retention to release']}", f"${fig['penalty per month']}",
                   f"${fig['surety line']}", f"${fig['bonds over the line']}",
                   "$29,091.69", "$8,125.00", "$85,235.00", "$61,477.00", "$38,915.00", "$25,080.00", "$17,204.84", "$7,132.50", "$13,040.00", "$18,360.00",
                   "$8,940.00", "$14,540.00", "$1,500.00", "$9,467.50", "November 9, 2026", "October 5, 2026", "September 28, 2026", "check 4460", "check 4474",
                   "check 4479", "joint check 4488", "May 9", "January 27", "certified mail number 7022 1670 0002 3391 4399", "1440 Eureka Road", "permit B25-1187", f"${fig['claims as recorded']}",
                   f"${fig['release demands amount']}"):
        rep.expect(f"note states {phrase}", phrase in ntext, True)
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
