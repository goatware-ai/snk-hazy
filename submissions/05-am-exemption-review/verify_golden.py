"""verify_golden.py for am-exemption-review: re-derive every figure the review states from inputs/ alone
under the company's policies as the golden applies them, and list the alternate readings a reviewer
could take with the phrase the golden uses to settle each one.

    .venv/bin/python drafts/05-am-exemption-review/verify_golden.py
"""
import csv
import datetime as dt
import os
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = os.environ.get("HAZY_ROOT") or next(str(p) for p in HERE.parents if (p / "tools" / "golden_verify.py").is_file())
sys.path.insert(0, os.path.join(ROOT, "tools"))
from golden_verify import Report  # noqa: E402

INPUTS = HERE / "inputs"
THRESHOLD = D("684.00"); STANDARD = D(40); MULT = D("1.5"); QWEEKS = 13


def xr(x, p=2):
    """Excel ROUND: half up off the 15 significant digit value of the float Excel holds; the overtime
    hours enter as the float difference of the badge hours and 40, as the workbook's MAX cell holds them."""
    return D(format(float(x), ".15g")).quantize(D(1).scaleb(-p), rounding=ROUND_HALF_UP)


def money(x):
    return f"{D(x):,.2f}"


def pdate(s):
    return dt.datetime.strptime(s, "%m/%d/%Y").date()


# badge file: three side-by-side blocks read back in block order
rows = list(csv.reader(open(INPUTS / "badge_hours_weekly_2024_2026.csv", newline="")))
hdr, body = rows[0], rows[1:]
WEEKS, HOURS = [], {"Voight": [], "Dellinger": [], "Natarajan": []}
for b in range(3):
    base = 4 * b
    for r in body:
        if r[base]:
            WEEKS.append(pdate(r[base]))
            HOURS["Voight"].append(D(r[base + 1])); HOURS["Dellinger"].append(D(r[base + 2])); HOURS["Natarajan"].append(D(r[base + 3]))
wb = openpyxl.load_workbook(INPUTS / "payroll_history_am.xlsx", data_only=True)
SAL = {"Voight": [], "Dellinger": [], "Natarajan": []}
for r in wb["Salary History"].iter_rows(min_row=5, values_only=True):
    if r[1]:
        SAL[r[1].split()[-1]].append((pdate(r[5]), D(str(r[6]))))
BON = {"Voight": {}, "Dellinger": {}, "Natarajan": {}}
QUARTERS = {}
for r in wb["Bonus History"].iter_rows(min_row=5, values_only=True):
    if r[1]:
        BON[r[1].split()[-1]][r[2]] = D(str(r[5])); QUARTERS[r[2]] = (pdate(r[3]), pdate(r[4]))
STAFF = {}
for r in wb["Store Staffing"].iter_rows(min_row=5, values_only=True):
    if isinstance(r[0], int):
        STAFF.setdefault(r[2].split()[-1], []).append((r[1], D(r[3]) + D(r[4]) / 2))
CURRENT = {k: v[-1][1] for k, v in SAL.items()}


def quarter_of(w):
    return next(q for q, (qs, qe) in QUARTERS.items() if qs <= w <= qe)


def derive(bonus_in_rate=True, pto_weeks_count=False, rate_denominator="40", half_time=D("0.5"), letter_rate=False):
    figures, standings = {}, {}
    for emp in ("Voight", "Dellinger", "Natarajan"):
        fte_now = STAFF[emp][-1][1] if half_time == D("0.5") else (D(STAFF[emp][-1][1] - D("0.5") * 0) if False else STAFF[emp][-1][1])
        fte_periods = [f for _, f in STAFF[emp]]
        salary_ok = CURRENT[emp] >= THRESHOLD
        duties_ok = sum(1 for f in fte_periods if f >= 2) > len(fte_periods) / 2
        exempt = salary_ok and duties_ok
        standings[f"{emp} classification"] = "Exempt" if exempt else "Non-exempt"
        ot_total = D(0); pay = D(0); worked = 0; over = 0; hrs = D(0)
        for k, w in enumerate(WEEKS):
            h = HOURS[emp][k]
            if h > 0 or pto_weeks_count:
                worked += 1; hrs += h
            ot = max(h - STANDARD, D(0))
            otf = max(float(h) - 40.0, 0.0)        # the workbook's MAX(hours-40,0) cell, a float difference
            if ot > 0: over += 1
            monday = w - dt.timedelta(days=6)
            sal = [s for e, s in SAL[emp] if e <= monday][-1]
            share = xr(float(BON[emp][quarter_of(w)]) / QWEEKS, 2) if bonus_in_rate else D(0)
            denom = STANDARD if rate_denominator == "40" else (h if h > 0 else STANDARD)
            rr = D("20.00") if letter_rate else xr((float(sal) + float(share)) / float(denom), 2)
            pay += xr(1.5 * float(rr) * otf, 2); ot_total += ot
        figures[f"{emp} workweeks worked"] = str(worked); figures[f"{emp} workweeks over 40"] = str(over)
        figures[f"{emp} overtime hours"] = f"{ot_total:.2f}"
        figures[f"{emp} back overtime"] = money(pay if not exempt else D(0))
        figures[f"{emp} overtime hours per workweek worked"] = f"{xr(ot_total / worked, 2)}"
        standings[f"{emp} back overtime"] = figures[f"{emp} back overtime"]
    company = sum(D(figures[f"{e} back overtime"].replace(",", "")) for e in ("Voight", "Dellinger", "Natarajan"))
    figures["company back overtime"] = money(company); figures["company exposure with liquidated damages"] = money(company * 2)
    v = D(figures["Voight back overtime"].replace(",", ""))
    figures["Voight exposure with liquidated damages"] = money(v * 2)
    figures["letter less review with liquidated damages"] = money(D("74880.00") - v * 2)
    q2_share = xr(float(BON["Voight"]["Q2 2025"]) / QWEEKS, 2)
    figures["Voight Q2 2025 bonus share"] = money(q2_share)
    figures["Voight Q2 2025 regular rate"] = money(xr((800.0 + float(q2_share)) / 40.0, 2))
    d_share = xr(float(BON["Dellinger"]["Q3 2025"]) / QWEEKS, 2)
    figures["Dellinger Q3 2025 regular rate"] = money(xr((675.0 + float(d_share)) / 40.0, 2))
    standings["company back overtime"] = figures["company back overtime"]
    return figures, standings


VARIANTS = {
    "bonuses left out of the regular rate": ({"bonus_in_rate": False}, "quarterly bonuses allocated across their thirteen workweeks"),
    "paid time off weeks counted as worked": ({"pto_weeks_count": True}, "weeks of paid time off at no hours"),
    "regular rate divided by hours worked": ({"rate_denominator": "hours"}, "divided by 40"),
    "letter's flat $20.00 rate applied": ({"letter_rate": True}, "our own policy puts them in"),
}

EXPECTED = {
    "Voight workweeks worked": "101", "Voight workweeks over 40": "101", "Voight overtime hours": "1054.08",
    "Voight back overtime": "32,685.45", "Voight overtime hours per workweek worked": "10.44",
    "Dellinger workweeks worked": "102", "Dellinger workweeks over 40": "82", "Dellinger overtime hours": "279.08",
    "Dellinger back overtime": "7,224.63",
    "Natarajan back overtime": "0.00",
    "company back overtime": "39,910.08", "company exposure with liquidated damages": "79,820.16",
    "Voight exposure with liquidated damages": "65,370.90", "letter less review with liquidated damages": "9,509.10",
    "Voight Q2 2025 bonus share": "41.54", "Voight Q2 2025 regular rate": "21.04", "Dellinger Q3 2025 regular rate": "17.73",
}

if __name__ == "__main__":
    rep = Report()
    figures, standings = derive()
    for label, golden in EXPECTED.items():
        rep.expect(label, figures.get(label), golden)
    rep.expect("Voight classification", standings["Voight classification"], "Non-exempt")
    rep.expect("Dellinger classification", standings["Dellinger classification"], "Non-exempt")
    rep.expect("Natarajan classification", standings["Natarajan classification"], "Exempt")
    rep.sensitivity(lambda **kw: derive(**kw)[1], VARIANTS)
    rep.finish()
